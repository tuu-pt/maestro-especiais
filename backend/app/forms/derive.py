"""Derive the empty form templates from the R1 forms in data/fixtures [A CONFIRMAR].

    python -m app.forms.derive  (make form-templates)

The ficha eletrotécnica loses the values of every mapped cell, the technician's block, the
cells filled by hand, the cached results of the formulas and the shared strings no cell uses.
The .docx forms lose every value cell of the numbered tables (1 to 4) and the date.
Reads only data/fixtures; never data/private.
"""

import io
import sys
from pathlib import Path

import docx

from app.forms.fill import (
    FE_BY_HAND,
    FE_SHEET,
    FE_TECHNICIAN,
    FE_VERSION,
    FORMS,
    clear_dates,
    labelled_pairs,
    section_of,
    unique_cells,
    write,
)  # fmt: skip
from app.forms.xlsx import Workbook
from app.ingest.ficha_eletrotecnica import cell_maps

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "data" / "fixtures" / "R1" / "MBERAL" / "2-PE" / "Editavel"
SOURCES = {
    "ficha_eletrotecnica": "MBERAL_FichaEletrotécnica_PE_ELE.xlsm",
    "identificacao": "MBERAL_Identificação do Projeto_PE_ELE.docx",
    "termo": "MBERAL_TermoResponsavilidade_PE_ELE.docx",
}
TEMPLATES = Path(__file__).parent / "templates"
# Rows of the tables with ticks, whose cells after the first are values
TICK_ROWS = ("SE/PS/PTC", "Rede MT/AT", "Rede BT", "Instalação de utilização MT/AT",
             "Instalação de utilização BT", "Grupos geradores")  # fmt: skip


def derive_fe(data: bytes) -> bytes:
    book = Workbook(data)
    part = book.sheet_part(FE_SHEET)
    refs = {*cell_maps()[FE_VERSION].cells, *FE_TECHNICIAN, *FE_BY_HAND}
    for ref in refs:
        book.put(part, ref, None)
    book.drop_hyperlinks(part, refs)
    for p in book.sheet_parts():
        book.clear_cached_results(p)
    book.prune_shared_strings()
    book.recalculate_on_open()
    return book.to_bytes()


def derive_docx(data: bytes) -> bytes:
    document = docx.Document(io.BytesIO(data))
    for table in document.tables:
        if not section_of(table).isdigit():
            continue
        for _, _, value in labelled_pairs(table):
            write(value, "")
        for row in table.rows:
            cells = unique_cells(row)
            if cells and " ".join(cells[0].text.split()) in TICK_ROWS:
                for c in cells[1:]:
                    write(c, "")
    clear_dates(document)
    core = document.core_properties
    core.author = core.last_modified_by = core.title = core.comments = ""
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def derive(source: Path = SOURCE) -> dict[str, bytes]:
    out = {}
    for kind, name in SOURCES.items():
        data = (source / name).read_bytes()
        out[kind] = derive_fe(data) if kind == "ficha_eletrotecnica" else derive_docx(data)
    return out


def main() -> None:
    TEMPLATES.mkdir(exist_ok=True)
    for kind, data in derive().items():
        (TEMPLATES / FORMS[kind].template).write_bytes(data)
        print(f"{FORMS[kind].template}: {len(data)} bytes", file=sys.stdout)


if __name__ == "__main__":
    main()
