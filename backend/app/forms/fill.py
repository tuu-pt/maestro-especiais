"""Pre-filled forms (SPEC 8.5): ficha eletrotécnica, Identificação and Termo de Responsabilidade.

What is written: the values of the confirmed ficha-base and the technician's profile. What never
is: the date and the signature (P8: the technician dates and signs), nor any tick the ficha-base
cannot justify. Whatever stays empty is listed, so the technician knows what is left.

The templates are the R1 forms with the values of the project cleared (app.forms.derive)
[A CONFIRMAR: until the TUU has its own empty templates].
"""

import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from importlib import resources
from typing import Any

import docx
from docx.table import _Cell

from app.assembly.values import ValueSource, format_value
from app.forms.xlsx import Workbook
from app.ingest.ficha_eletrotecnica import cell_maps
from app.ingest.keys import KEYS
from app.library.facts import DOC_KEYS

FE_VERSION = "FE_v.20190222"
FE_SHEET = "Ficha Eletrotecnica"
# The technician's block of the ficha eletrotécnica (not read: it comes from the profile).
FE_TECHNICIAN = {"C11": "tec.nome", "Q11": "tec.nif", "C12": "tec.telefone", "J12": "tec.email",
                 "Q12": "tec.dgeg"}  # fmt: skip
# Filled by the technician: date (M40), requerente's phone (C6, no key in SPEC 7.2) and the
# columns of the entry with no key (entrada, ramal, andar, fração).
FE_BY_HAND = {"C6": "Telefone do requerente", "C29": "Entrada do imóvel", "D29": "Ramal n.º",
              "K29": "Andar", "L29": "Fração", "M40": "Data (pelo técnico)"}  # fmt: skip

PROMOTOR = {"Nome": "id.requerente.nome", "Telefone": None, "E-mail": "id.requerente.email",
            "NIF": "id.requerente.nif", "Morada": "id.requerente.morada",
            "C. Postal": "id.requerente.cp"}  # fmt: skip
TECHNICIAN = {"Nome": "tec.nome", "N.º BI/CC": "tec.cc", "Telefone": "tec.telefone",
              "E-mail": "tec.email", "NIF": "tec.nif", "N.º DGEG": "tec.dgeg", "N.º OE": None,
              "N.º OET": "tec.oet", "Morada": "tec.morada", "C. Postal": "tec.cp"}  # fmt: skip
PLACE = {"Lugar/Rua": "id.local.rua", "Freguesia": "id.local.freguesia",
         "Concelho": "id.local.concelho", "Distrito": "id.local.distrito"}  # fmt: skip
# Labels of the DGEG forms, per numbered section → ficha-base key (None: by hand) [A CONFIRMAR]
DOCX_FIELDS: dict[str, dict[str, dict[str, str | None]]] = {
    "identificacao": {
        "1": PROMOTOR,
        "2": TECHNICIAN,
        "3": {
            **PLACE,
            "Coordenadas GPS": "id.local.gps",
            "NIP": "id.local.nip",
            "Tipo de estabelecimento": "ele.tipo_utilizacao",
            "Tensão da RESP [kV]": "ele.tensao_resp_kv",
            "Potência a alimentar pela RESP [kVA]": "ele.potencia_alimentar_kva",
        },
    },
    "termo": {
        "1": {k: v for k, v in PROMOTOR.items() if k in ("Nome", "Telefone", "E-mail", "NIF")},
        "2": TECHNICIAN,
        "3": {**PLACE, "Tipo de estabelecimento": "ele.classificacao"},
        "4": {"NIP": "id.local.nip", "CPE(s)": None},
    },
}
# Ticks: (section, label of the row) → (key, value that ticks it)
DOCX_TICKS: dict[str, dict[tuple[str, str], tuple[str, str]]] = {
    "identificacao": {},  # section 4: by the technician [A CONFIRMAR: which row the ficha implies]
    "termo": {
        ("4", "Instalação nova"): ("ele.instalacao", "Nova"),
        ("4", "Instalação existente"): ("ele.instalacao", "Existente"),
    },
}
TICK_LABELS = ("Instalação nova", "Instalação existente")
DATE = re.compile(r"^\s*\d{1,4}[/.-]\d{1,2}[/.-]\d{1,4}\s*$")


@dataclass(frozen=True)
class FormKind:
    kind: str
    title: str
    template: str
    suffix: str
    media_type: str


XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FORMS = {
    "ficha_eletrotecnica": FormKind("ficha_eletrotecnica", "Ficha eletrotécnica",
                                    "fe_v20190222.xlsm", "FichaEletrotecnica.xlsm", XLSM),
    "identificacao": FormKind("identificacao", "Identificação do Projeto",
                              "identificacao_v2018_1.docx", "IdentificacaoProjeto.docx", DOCX),
    "termo": FormKind("termo", "Termo de Responsabilidade",
                      "termo_v2018_1.docx", "TermoResponsabilidade.docx", DOCX),
}  # fmt: skip


@dataclass
class Filled:
    data: bytes
    by_hand: list[str] = field(default_factory=list)  # what the technician still fills


def template(name: str) -> bytes:
    return resources.files("app.forms").joinpath("templates", name).read_bytes()


def label(key: str) -> str:
    info = KEYS.get(key)
    return info.label_pt if info else DOC_KEYS.get(key, key)


def _raw(values: ValueSource, key: str) -> Any:
    if key.startswith(("tec.", "doc.")):
        return values.profile.get(key) or None
    fv = values.values.get(key)
    return None if fv is None or fv.value in ("", []) else fv.value


def _cell_value(value: Any) -> Any:
    """What goes into a cell: numbers stay numbers (the FE formulas use them)."""
    if isinstance(value, list):
        return ", ".join(format_value(v) for v in value)
    return value


def fill_fe(values: ValueSource) -> Filled:
    book = Workbook(template(FORMS["ficha_eletrotecnica"].template))
    part = book.sheet_part(FE_SHEET)
    cells = {**cell_maps()[FE_VERSION].cells, **FE_TECHNICIAN}
    missing = []
    for ref, key in cells.items():
        value = _raw(values, key)
        book.put(part, ref, _cell_value(value))
        if value is None:
            missing.append(f"{label(key)} ({ref})")
    book.recalculate_on_open()
    return Filled(book.to_bytes(), missing + [f"{v} ({k})" for k, v in FE_BY_HAND.items()])


# ---------------------------------------------------------------- docx


def _text(cell: _Cell) -> str:
    return " ".join(cell.text.split())


def _label(text: str) -> str:
    return text.rstrip(":").strip()


def unique_cells(row: Any) -> list[_Cell]:
    """The cells of a row once each (python-docx repeats merged cells)."""
    seen: list[Any] = []
    out = []
    for c in row.cells:
        if c._tc not in seen:
            seen.append(c._tc)
            out.append(c)
    return out


def section_of(table: Any) -> str:
    first = unique_cells(table.rows[0])
    return _text(first[0]) if first else ""


def write(cell: _Cell, text: str) -> None:
    """Replace a cell's text, keeping the first run's format."""
    paragraphs = cell.paragraphs
    for p in paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = paragraphs[0]
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = text
    elif text:
        p.add_run(text)


def labelled_pairs(table: Any) -> list[tuple[str, _Cell, _Cell]]:
    """(label, label cell, value cell) for each "Label:" cell of a table."""
    pairs = []
    for row in table.rows[1:]:
        cells = unique_cells(row)
        labelled_row = bool(cells) and _text(cells[0]).endswith(":")  # not a header row
        for i, c in enumerate(cells[:-1]):
            text = _text(c)
            if text.endswith(":") or (labelled_row and text in TICK_LABELS):
                pairs.append((_label(text), c, cells[i + 1]))
    return pairs


def clear_dates(document: Any) -> None:
    """P8: the date is the technician's. Paragraphs with only a date become empty."""
    for table in document.tables:
        for row in table.rows:
            for cell in unique_cells(row):
                for p in cell.paragraphs:
                    if DATE.match(p.text):
                        for r in p.runs:
                            r.text = ""


def fill_docx(kind: str, values: ValueSource) -> Filled:
    document = docx.Document(io.BytesIO(template(FORMS[kind].template)))
    fields, ticks = DOCX_FIELDS[kind], DOCX_TICKS[kind]
    by_hand: list[str] = []
    for table in document.tables:
        section = section_of(table)
        for name, _, cell in labelled_pairs(table):
            if (section, name) in ticks:
                key, wanted = ticks[(section, name)]
                raw = _raw(values, key)
                write(cell, "X" if raw is not None and str(raw).strip() == wanted else "")
                if raw is None:
                    by_hand.append(f"{section}. {name}")
                continue
            if name not in fields.get(section, {}):
                continue
            key_or_none = fields[section][name]
            raw = _raw(values, key_or_none) if key_or_none else None
            write(cell, "" if raw is None else format_value(raw))
            if raw is None:
                by_hand.append(f"{section}. {name}")
    if kind == "identificacao":
        by_hand.append("4. Tipo de instalação (nova ou existente, por linha)")
    by_hand.append("Data e assinatura do técnico responsável")
    clear_dates(document)
    out = io.BytesIO()
    document.save(out)
    return Filled(out.getvalue(), by_hand)


FILLERS: dict[str, Callable[[ValueSource], Filled]] = {
    "ficha_eletrotecnica": fill_fe,
    "identificacao": lambda v: fill_docx("identificacao", v),
    "termo": lambda v: fill_docx("termo", v),
}


def fill(kind: str, values: ValueSource) -> Filled:
    return FILLERS[kind](values)
