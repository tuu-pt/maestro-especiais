"""Reader of the MQT (mapa de quantidades) and the LPU (lista de preços unitários), .xlsx.

Two header variants (SPEC 8.2): "CÓDIGO | DESIGNAÇÃO | UNI. | QUANT. | PREÇO…" (MQT of R1) and
"Artº | Designação | Un | QUANTIDADES ADJUDICADAS" with a second line "Quant. | Pr. Unit. |
Total | Total Cap." (LPU of R2). Columns are found by their header text. Every line becomes a
BomItem: chapter, subchapter, article (has a unit or a quantity), description, note or total.
The identification above the header ("Designação:", "Adjudicante:") goes to the ficha-base,
where a divergence is a FichaConflict (case C7). Nothing is computed (P2).
"""

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.ingest import consolidate
from app.ingest.base import Candidate, ReadResult, to_number
from app.ingest.detect import fold
from app.ingest.pipeline import ReaderError
from app.models import BomItem, ProjectFile

COLUMNS: tuple[tuple[str, str], ...] = (
    ("code", r"^(codigo|art\.?o?|art\.?º|artigo|ref)"),
    ("designation", r"^designa"),
    ("unit", r"^(uni|un)\b"),
    ("chapter_total", r"total\s*cap"),
    ("unit_price", r"pre[cç]o\s*unit|pr\.?\s*unit"),
    ("total", r"^(pre[cç]o\s*)?total"),
    ("quantity", r"^quant"),
)
# Identification rows above the header: label → ficha key (SPEC 7.2).
ID_LABELS = {"designacao": "id.obra.designacao", "adjudicante": "id.requerente.nome"}
_CODE = re.compile(r"^\d+(?:\.\d+)*$")
_BOARD = re.compile(r"^Q[.\s]?[A-Z0-9]", re.I)


@dataclass
class BomLine:
    row_index: int
    source_ref: str
    kind: str
    code: str | None = None
    level: int = 0
    parent_code: str | None = None
    designation: str | None = None
    unit: str | None = None
    quantity: Decimal | None = None
    unit_price: Decimal | None = None
    total: Decimal | None = None
    chapter_total: Decimal | None = None


@dataclass
class BomReading:
    variant: str  # "mqt" | "lpu"
    lines: list[BomLine]
    result: ReadResult
    warnings: list[str] = field(default_factory=list)


def variant_of(rows: list[tuple[Any, ...]]) -> str | None:
    """By the title of the document, never by the file name."""
    for row in rows[:12]:
        text = " ".join(fold(c) for c in row if c is not None)
        if "lista de precos" in text:
            return "lpu"
        if "mapa de quantidade" in text or "mapa de trabalhos" in text:
            return "mqt"
    return None


def _find_header(rows: list[tuple[Any, ...]]) -> int | None:
    for index, row in enumerate(rows[:40]):
        words = [fold(c) for c in row if c is not None]
        if any(w.startswith("designa") for w in words) and any(
            w.startswith("quant") for w in words
        ):
            return index
    return None


def _columns(rows: list[tuple[Any, ...]], index: int) -> tuple[dict[int, str], int, list[str]]:
    """Column → field, first data row, and the headers not understood."""
    below = rows[index + 1] if index + 1 < len(rows) else ()
    texts = [c for c in below if c is not None and str(c).strip()]
    # A second line only if it names columns ("Quant. | Pr. Unit. | Total"), not a first row.
    two_lines = (
        bool(texts)
        and all(to_number(t) is None for t in texts)
        and any(re.search(p, fold(t)) for t in texts for _, p in COLUMNS)
    )
    columns: dict[int, str] = {}
    unknown: list[str] = []
    width = max(len(rows[index]), len(below))
    for i in range(width):
        headers = [rows[index][i] if i < len(rows[index]) else None]
        if two_lines and i < len(below):
            headers.insert(0, below[i])  # the second line names the column when it has text
        for header in headers:
            if header is None or not str(header).strip():
                continue
            name = next((f for f, p in COLUMNS if re.search(p, fold(header))), None)
            if name and name not in columns.values():
                columns[i] = name
                break
            if name is None and "quantidades" not in fold(header):
                unknown.append(" ".join(str(header).split())[:40])
    return columns, index + (2 if two_lines else 1), unknown


def _identification(rows: list[tuple[Any, ...]], header: int, sheet: str) -> list[Candidate]:
    out: list[Candidate] = []
    for r, row in enumerate(rows[:header], start=1):
        cells = [(i, c) for i, c in enumerate(row) if c is not None and str(c).strip()]
        for (_, label), nxt in zip(cells, [*cells[1:], (None, None)], strict=False):
            key = ID_LABELS.get(fold(label).rstrip(":").strip())
            if key and nxt[1] is not None and not str(nxt[1]).endswith(":"):
                col = get_column_letter((nxt[0] or 0) + 1)
                out.append(Candidate(key, str(nxt[1]).strip(), f"{sheet}!{col}{r}"))
    return out


def _board_name(designation: str) -> str | None:
    name = re.sub(r"\s*\(.*?\)\s*$", "", designation).strip()
    return name if _BOARD.match(name) and len(name) <= 30 else None


def read(data: bytes) -> BomReading:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise ReaderError("MQT/LPU ilegível ou corrompido.") from exc
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            header = _find_header(rows)
            if header is not None:
                return _read_sheet(sheet.title, rows, header)
    finally:
        workbook.close()
    raise ReaderError(
        "Cabeçalho do MQT/LPU não reconhecido (DESIGNAÇÃO e QUANT.): é preciso rever o ficheiro."
    )


def _read_sheet(title: str, rows: list[tuple[Any, ...]], header: int) -> BomReading:
    columns, first, unknown = _columns(rows, header)
    variant = variant_of(rows) or "mqt"
    result = ReadResult(source_type="mqt")
    reading = BomReading(variant=variant, lines=[], result=result)
    if variant_of(rows) is None:
        reading.warnings.append("Título do documento não encontrado: tratado como MQT.")
    reading.warnings += [f"Coluna não reconhecida: «{u}»." for u in unknown if u]
    if "designation" not in columns.values():
        raise ReaderError("Sem coluna de designação no MQT/LPU.")
    result.values = _identification(rows, header, title)
    boards: list[str] = []
    for r, row in enumerate(rows[first:], start=first + 1):
        cell = {name: row[i] for i, name in columns.items() if i < len(row)}
        texts = [c for c in row if c is not None and str(c).strip()]
        if not texts:
            continue
        line = _line(title, r, cell, texts)
        reading.lines.append(line)
        if line.kind == "article" and line.designation:
            board = _board_name(line.designation)
            if board and board not in boards:
                boards.append(board)
    if boards:
        result.values.append(Candidate("ele.quadros", boards, f"{title} · artigos de quadros"))
    if not any(line.kind == "article" for line in reading.lines):
        reading.warnings.append("O MQT/LPU não tem artigos com unidade ou quantidade.")
    return reading


def _line(title: str, r: int, cell: dict[str, Any], texts: list[Any]) -> BomLine:
    raw_code = str(cell.get("code") or "").strip()
    code = raw_code if _CODE.match(raw_code) else None
    designation = str(cell["designation"]).strip() if cell.get("designation") else None
    unit = str(cell["unit"]).strip() if cell.get("unit") else None
    line = BomLine(
        row_index=r,
        source_ref=f"{title}!linha {r}",
        kind="description",
        code=code,
        level=code.count(".") + 1 if code else 0,
        parent_code=code.rsplit(".", 1)[0] if code and "." in code else None,
        designation=designation,
        unit=unit,
        quantity=to_number(cell.get("quantity")),
        unit_price=to_number(cell.get("unit_price")),
        total=to_number(cell.get("total")),
        chapter_total=to_number(cell.get("chapter_total")),
    )
    folded = fold(designation or " ".join(str(t) for t in texts))
    if line.unit or line.quantity is not None:
        line.kind = "article"
    elif code and line.level == 1:
        line.kind = "chapter"
    elif code and line.level == 2:
        line.kind = "subchapter"
    elif not code and folded.startswith("nota"):
        line.kind = "note"
    elif not code and ("total" in folded or designation is None):
        line.kind = "total"
    return line


def add_reading(db: Session, file: ProjectFile, reading: BomReading) -> str:
    """Store the lines (replacing the previous file of the same variant) and merge the values."""
    revision = consolidate.draft_revision(db, file.project_id, None)
    db.execute(
        delete(BomItem).where(
            BomItem.revision_id == revision.id, BomItem.variant == reading.variant
        )
    )
    for line in reading.lines:
        db.add(
            BomItem(
                revision_id=revision.id, source_file_id=file.id, variant=reading.variant,
                row_index=line.row_index, source_ref=line.source_ref, code=line.code,
                level=line.level, parent_code=line.parent_code, kind=line.kind,
                designation=line.designation, unit=line.unit, quantity=line.quantity,
                unit_price=line.unit_price, total=line.total, chapter_total=line.chapter_total,
                link_status="unlinked",
            )
        )  # fmt: skip
    reading.result.warnings = reading.warnings
    summary = consolidate.apply(db, file, reading.result)
    articles = sum(1 for line in reading.lines if line.kind == "article")
    label = "LPU" if reading.variant == "lpu" else "MQT"
    return f"{label}: {articles} artigos · {summary}"
