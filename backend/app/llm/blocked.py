"""Names the privacy guard blocks, seeded in development from the reference projects (P9).

The TUU team and technicians: the names of the title blocks of the drawings ("NOME | ENG …"),
of the technician in the signatures of the MDJ/CTE and in the forms. Some of these were kept in
the fixtures on purpose (decision of 24 set 2026); they must never reach the LLM. In production
the list comes from the technicians' profiles and from what an admin adds.
"""

import re
from collections.abc import Iterator
from pathlib import Path

import docx
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ingest.detect import fold
from app.library.docx_blocks import split
from app.library.facts import signature_facts
from app.library.sources import reference_documents
from app.models import BlockedTerm

_TITLE_BLOCK = re.compile(
    r"([A-ZÀ-Ý][A-Za-zÀ-ÿ'-]+(?:[ \t]+[A-ZÀ-Ý][A-Za-zÀ-ÿ'-]+){1,4})[ \t]*\|[ \t]*ENG", re.I
)
_NAME = re.compile(r"^[A-Za-zÀ-ÿ'-]+(?: [A-Za-zÀ-ÿ'-]+){1,4}$")
MAX_PAGES = 40


def _clean(name: str) -> str | None:
    name = " ".join(name.split()).strip(" .,;:")
    return name if _NAME.match(name) and 5 <= len(name) <= 60 else None


def _drawing_names(path: Path) -> Iterator[str]:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        for i in range(min(len(document), MAX_PAGES)):
            text = document[i].get_textpage().get_text_range()
            for line in text.splitlines():
                for m in _TITLE_BLOCK.finditer(line):
                    if name := _clean(m.group(1)):
                        yield name
    finally:
        document.close()


def _form_names(path: Path) -> Iterator[str]:
    """The technician of the forms (identificação, termo): "Nome:" in the technician's table."""
    document = docx.Document(str(path))
    for table in document.tables:
        text = fold(" ".join(c.text for c in table.rows[0].cells)) if table.rows else ""
        if "tecnico responsavel" not in text:
            continue
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if cells and fold(cells[0]).startswith("nome"):
                value = next((c for c in cells[1:] if c and c != cells[0]), "")
                if name := _clean(value):
                    yield name


def names_in(fixtures: Path, code: str) -> set[str]:
    names: set[str] = set()
    root = fixtures / code
    for path in sorted(root.rglob("*.pdf")):
        names |= set(_drawing_names(path))
    for path in sorted(root.rglob("*.docx")):
        if "~$" not in path.name:
            names |= set(_form_names(path))
    for _, _, data in reference_documents(fixtures, code):
        parts = split(data)
        signature = next((s for s in parts.sections if s.kind == "signature"), None)
        if signature:
            names |= {f.value for f in signature_facts(signature.text.splitlines())
                      if f.key == "tec.nome" and _clean(f.value)}  # fmt: skip
    return names


def seed_blocked_terms(db: Session, fixtures: Path, codes: tuple[str, ...]) -> int:
    known = {fold(t.value) for t in db.scalars(select(BlockedTerm))}
    added = 0
    for code in codes:
        if not (fixtures / code).is_dir():
            continue
        for name in sorted(names_in(fixtures, code)):
            if fold(name) not in known:
                db.add(BlockedTerm(value=name, kind="name", source="fixtures"))
                known.add(fold(name))
                added += 1
    db.flush()
    return len(known)
