"""Synthetic project files for tests only (SPEC 8.2 structures). Invented values, no real data.

The acceptance tests with the anonymized R1/R2 fixtures live in tests/acceptance/.
"""

import io
from typing import Any

import openpyxl
import xlwt

FE_VERSION = "FE_v.20190222"

# A ficha eletrotécnica laid out on the fixed cells of FE_v.20190222.
FE_DEFAULT: dict[str, Any] = {
    # The cell to the right of each label of the DGEG template (maps/fe_v20190222.yaml).
    "C5": "Requerente Sintético",
    "Q5": 999990013,
    "J6": "requerente@example.com",
    "C7": "Rua de Teste 1",
    "C8": "0000-001 Localidade de Teste",
    "C15": "Freguesia de Teste",
    "M15": "Concelho de Teste",
    "Q15": "Distrito de Teste",
    "G16": "Rua de Teste",
    "Q16": "0.000, -0.000",
    "E29": "NIP-TESTE-1",
    "F23": "Unifamiliar",
    "Q23": "Nova",
    "F24": "Locais de habitação",
    "Q24": 1,
    "B29": "C",
    "M29": "Habitação",
    "O29": "Trif",
    "P29": 34.5,
    "Q29": 1,
    "R29": 34.5,
    "C11": "Técnico Sintético",
    "Q12": "000123",
    "I44": 34.5,
    "R45": FE_VERSION,
}


def workbook_bytes(sheets: dict[str, list[list[Any]] | dict[str, Any]]) -> bytes:
    """A workbook whose sheets are given as rows (list) or as cell → value (dict)."""
    wb = openpyxl.Workbook()
    first = True
    for title, content in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = title
        first = False
        if isinstance(content, dict):
            for cell, value in content.items():
                ws[cell] = value
        else:
            for row in content:
                ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def ficha_eletrotecnica(**cells: Any) -> bytes:
    values = {**FE_DEFAULT, **cells}
    return workbook_bytes(
        {"Ficha Eletrotecnica": {k: v for k, v in values.items() if v is not None}}
    )


CALC_HEADER = [
    "ORIGEM", "DESTINO", "Potência (kVA)", "Tensão (V)", "Tipo de proteção", "IB (A)", "In (A)",
    "IΔn (mA)", "Iz (A)", "I2 (A)", "1,45 x Iz (A)", "Cabo", "Comprimento (m)",
    "QDT troço (%)", "QDT montante (%)", "QDT total (%)", "PdC (kA)", "Pólos", "Instalação",
    "Fases", "Isolamento", "Condutor", "Método de referência", "Quadro RTIEBT",
]  # fmt: skip

CALC_ROWS: list[list[Any]] = [
    ["ENTRADA DE ENERGIA"],
    ["Portinhola", "Q.E.G.", 34.5, 400, "F", 50, 63, None, 76, 100.8, 110.2, "XZ1(frt,zh) 4x16",
     25, 0.8, 0, 0.8, 6, "MUL", "ENT", 3, "XLPE", "Cu", "D1", "52-C3"],
    ["EDIFÍCIO"],
    ["Q.E.G.", "Q.P.1", 12, 400, "D", 18, 25, 30, 32, 36.25, 46.4, "H07V-U 3G6",
     18, 1.1, 0.8, 1.9, 6, "MUL", "TUB", 3, "PVC", "Cu", "B1", "52-C1"],
    ["Q.E.G.", "C1 Iluminação", 1.2, 230, "D", 5.2, 10, 30, 15.5, 14.5, 22.5, "H07V-U 3G1,5",
     20, 1.4, 0.8, 2.2, 6, "MON", "TUB", 1, "PVC", "Cu", "B1", "52-C1"],
]  # fmt: skip


def tabela_calculo(header: list[str] | None = None, rows: list[list[Any]] | None = None) -> bytes:
    return workbook_bytes({"Tabela": [header or CALC_HEADER, *(rows or CALC_ROWS)]})


# ---------------------------------------------------------------- 09-Folha de Cálculo (.xls)

FOLHA09_TITLES = {
    "IB": "CÁLCULO CORRENTE DE SERVIÇO IB",
    "condutores": "SECÇÃO DOS CONDUTORES OU CABO",
    "tensao": "Queda de Tensão",
    "proteccao": "Dimensionamento da Protecção",
}
# Q.E.G. → Q.P.1 of CALC_ROWS, as the TUU sheet writes it (voltage drop as a fraction).
FOLHA09_VALUES: dict[str, Any] = {
    "IB!H7": 12,
    "IB!H17": 18.0412,
    "condutores!N8": 18,
    "condutores!H47": 6,
    "proteccao!Z7": 32,
    "proteccao!E9": 25,
    "proteccao!J9": 36.25,
    "proteccao!U15": 46.4,
    "tensao!T12": 0.011,
}


def _xls_cell(ref: str) -> tuple[str, int, int]:
    sheet, cell = ref.split("!")
    letters = "".join(c for c in cell if c.isalpha())
    column = 0
    for c in letters:
        column = column * 26 + ord(c) - 64
    return sheet, int(cell[len(letters) :]) - 1, column - 1


def folha09(
    titles: dict[str, str] | None = None, sheets: tuple[str, ...] | None = None, **cells: Any
) -> bytes:
    """A 09-Folha with the TUU layout; cell overrides as IB_H7=…, None to leave a cell empty."""
    values = {**FOLHA09_VALUES, **{k.replace("_", "!", 1): v for k, v in cells.items()}}
    book = xlwt.Workbook()
    names = sheets or ("Dimensionamento", "IB", "condutores", "tensao", "proteccao", "impressao")
    ws = {name: book.add_sheet(name) for name in names}
    for name, title in (titles or FOLHA09_TITLES).items():
        if name in ws:
            ws[name].write(1, 1, title)
    for ref, value in values.items():
        sheet, row, col = _xls_cell(ref)
        if sheet in ws and value is not None:
            ws[sheet].write(row, col, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def mqt() -> bytes:
    return workbook_bytes(
        {"MQT": [["CÓDIGO", "DESIGNAÇÃO", "UNI.", "QUANT."], ["1.1", "Quadro Q.E.G.", "un", 1]]}
    )


# ---------------------------------------------------------------- MQT / LPU (.xlsx)

MQT_ROWS: list[list[Any]] = [
    [None, "MORADIA DE TESTE"],
    [None, "MAPA DE QUANTIDADE DE TRABALHOS - ELETRICIDADE"],
    [],
    [None, "CÓDIGO", "DESIGNAÇÃO", "UNI.", "QUANT.", "PREÇO UNITÁRIO", "PREÇO TOTAL",
     "TOTAL CAPÍTULO"],
    [None, None, "NOTAS INICIAIS"],
    [None, None, "Nota 1. Os preços não incluem IVA."],
    [None, "8", "ELETRICIDADE", None, None, None, None, 4300],
    [None, "8.1", "ENTRADA E DISTRIBUIÇÃO DE ENERGIA"],
    [None, "8.1.1", "Caixa Portinhola P100", "cj", 1, 250, 250],
    [None, "8.2", "QUADROS ELÉTRICOS"],
    [None, "8.2.1", "Fornecimento e instalação dos quadros elétricos"],
    [None, "8.2.1.1", "Q.E.G", "un", 1, 1100, 1100],
    [None, "8.2.1.2", "Q.P.1", "un", 1, 600, 600],
    [None, "8.3", "CARREGAMENTO DE VEÍCULOS ELÉTRICOS"],
    [None, "8.3.1", "Carregador de veículos elétricos 7,4 kW", "un", 2, 1175, 2350],
    [None, None, None, None, None, None, "TOTAL:", 4300],
]  # fmt: skip
LPU_ROWS: list[list[Any]] = [
    [None, "LISTA DE PREÇOS UNITÁRIOS"],
    [None, "Designação:", "Reabilitação do Edifício de Teste"],
    [None, "Adjudicatário:"],
    [None, "Adjudicante:", "Município de Teste"],
    [None, "Data:", "2026-01-15"],
    [],
    [None, "Artº", "Designação", "Un", "QUANTIDADES ADJUDICADAS"],
    [None, None, None, None, "Quant.", "Pr. Unit.", "Total", "Total Cap."],
    [None, "1", "INSTALAÇÕES ELÉTRICAS", None, None, None, None, 7450],
    [None, "1.8", "Quadros elétricos"],
    [None, "1.8.1", "Fornecimento e montagem de quadros elétricos"],
    [None, "1.8.1.1", "Q.E.G.", "Un", 1, 6470, 6470],
    [None, "1.8.1.2", "Q.P.1 (piso 1)", "Un", 1, 980, 980],
]  # fmt: skip


def bom(rows: list[list[Any]], title: str = "MQT") -> bytes:
    return workbook_bytes({title: rows})


PDF_MINIMAL = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n"
