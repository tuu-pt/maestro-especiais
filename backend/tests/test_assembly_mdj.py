"""MDJ of R1 assembled from its confirmed ficha-base (Phase 4, task 1)."""

import io
import re
import zipfile
from typing import Any

import docx
import pytest
from conftest import Api
from lxml import etree
from reference_projects import load_confirmed, seed_library
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.library.docx_blocks import DOCUMENT, w
from app.models import Section, TemplateBlock, ValueRef
from app.storage import ObjectStore

pytestmark = pytest.mark.usefixtures("inline_ingestion")


@pytest.fixture
def mdj(api: Api, db: Session, store: ObjectStore) -> dict[str, Any]:
    seed_library(db, store)
    project_id = load_confirmed(api, "R1")
    response = api.as_("redator").post(
        f"/api/projects/{project_id}/documents", json={"type": "MDJ"}
    )
    assert response.status_code == 201, response.text
    return dict(response.json())


def section(doc: dict[str, Any], suffix: str) -> dict[str, Any]:
    return next(s for s in doc["sections"] if s["block_key"].endswith(suffix))


def texts(s: dict[str, Any]) -> list[str]:
    out = []
    for node in s["content"]["content"]:
        children = node.get("content") or []
        if node["type"] == "locked":
            children = [c for p in children for c in p.get("content") or []]
        out.append("".join(c.get("text", "") for c in children))
    return out


def _root(data: bytes) -> Any:
    return etree.fromstring(zipfile.ZipFile(io.BytesIO(data)).read(DOCUMENT))


def _fragment(data: bytes, fragment: str) -> list[Any]:
    root = _root(data)
    decl = " ".join(f'xmlns:{p}="{u}"' for p, u in root.nsmap.items() if p)
    return list(etree.fromstring(f"<wrap {decl}>{fragment}</wrap>"))


def _c14n(el: Any) -> bytes:
    return bytes(etree.tostring(el, method="c14n"))


def test_only_from_a_confirmed_ficha_base(api: Api, db: Session, store: ObjectStore) -> None:
    seed_library(db, store)
    project = api.as_("redator").post("/api/projects", json={"code": "X1", "name": "X"}).json()

    response = api.as_("redator").post(
        f"/api/projects/{project['id']}/documents", json={"type": "MDJ"}
    )

    assert response.status_code == 409
    assert "ficha-base ainda não foi confirmada" in response.json()["detail"]


def test_sections_follow_the_library_with_the_rules_evaluated(mdj: dict[str, Any]) -> None:
    assert len(mdj["sections"]) == 42
    assert [s["order"] for s in mdj["sections"]] == list(range(1, 43))
    rpc = section(mdj, "regulamento_dos_produtos_de_construcao_rpc")
    assert not rpc["active"]
    assert rpc["active_reason"] == (
        "«Classificação do local» é «Locais de habitação» (a regra pede != «Locais de habitação»)."
    )
    pv = section(mdj, "instalacao_fotovoltaica")
    assert not pv["active"] and "Fotovoltaico" in pv["active_reason"]
    buried = section(mdj, "canalizacoes.canalizacoes_enterradas")
    assert buried["active"]  # the trench of the MQT of R1
    assert all(not s["block_approved"] for s in mdj["sections"])  # every block is still proposed
    assert mdj["counts"]["not_approved"] == 42


def test_power_is_resolved_from_the_ficha_with_its_value_ref(
    mdj: dict[str, Any], db: Session
) -> None:
    supply = section(mdj, "alimentacao_de_energia")

    assert any("potência elétrica de 34,5 kVA" in t for t in texts(supply))
    refs = db.scalars(select(ValueRef).where(ValueRef.key == "ele.potencia_alimentar_kva")).all()
    assert refs and refs[0].rendered_text == "34,5" and refs[0].ficha_value_id is not None
    # the adaptive paragraphs wait for the drafting step
    assert supply["status"] == "todo" and supply["status_note"] == "Por gerar: texto adaptativo."


def test_fixed_blocks_are_locked_with_the_library_ooxml(mdj: dict[str, Any], db: Session) -> None:
    s = section(mdj, "dimensionamento_eletrico.quedas_de_tensao")

    assert s["mode"] == "fixed" and s["locked"] and s["status"] == "generated"
    assert all(n["type"] == "locked" for n in s["content"]["content"])
    ip_ik = section(mdj, "caracteristicas_dos_equipamentos_em_funcao_das_influencias_externas")
    assert ip_ik["locked"] and ip_ik["status"] == "generated"


def test_a_skeleton_block_with_no_text_is_todo(mdj: dict[str, Any]) -> None:
    lighting = section(mdj, "instalacoes_eletricas_a_considerar.iluminacao_de_seguranca")
    assert lighting["active"] and lighting["status"] == "todo"
    assert lighting["status_note"].startswith("Sem texto: nenhum projeto de referência")


def test_the_signature_comes_from_the_profile_of_the_technician(mdj: dict[str, Any]) -> None:
    # R1 was confirmed by dev:tecnico: with DEV_AUTH, the fake development profile
    signature = section(mdj, "ele.mdj.assinatura")
    text = " ".join(texts(signature))

    assert signature["status"] == "generated" and not signature["missing_keys"]
    assert "Localidade (desenvolvimento)" in text  # doc.local is not personal
    assert "Técnico de Desenvolvimento" not in text and "•••" in text  # masked in the API
    assert "[data: pelo técnico]" in text  # P8, never missing


def test_missing_values_leave_the_section_todo(mdj: dict[str, Any]) -> None:
    cover = section(mdj, "ele.mdj.capa")

    assert cover["status"] == "todo"
    assert cover["status_note"].startswith("Falta dado: ")
    # the cover: requerente masked, obra missing in the ficha eletrotécnica of R1
    assert "REQUERENTE: •••" in texts(cover)
    assert "id.obra.designacao" in cover["missing_keys"]


def test_personal_values_are_never_stored_in_the_content(mdj: dict[str, Any], db: Session) -> None:
    refs = db.scalars(select(ValueRef).where(ValueRef.personal.is_(True))).all()
    assert refs and all(r.rendered_text is None for r in refs)
    assert "Bruno Exemplo" not in str(mdj)


def test_export_check_lists_what_stops_the_official_export(api: Api, mdj: dict[str, Any]) -> None:
    check = api.as_("redator").get(f"/api/documents/{mdj['id']}/export-check").json()

    assert check["ready"] is False
    reasons = {p["reason"] for p in check["problems"]}
    assert "bloco não aprovado" in reasons and "Por gerar: texto adaptativo." in reasons


def test_draft_docx_keeps_the_fixed_ooxml(api: Api, mdj: dict[str, Any], db: Session) -> None:
    response = api.as_("redator").get(f"/api/documents/{mdj['id']}/draft.docx")

    assert response.status_code == 200
    assert "R1_MDJ_rascunho.docx" in response.headers["content-disposition"]
    data = response.content
    document = docx.Document(io.BytesIO(data))
    body = "\n".join(p.text for p in document.paragraphs)
    assert "34,5 kVA" in body
    assert "REGULAMENTO DOS PRODUTOS" not in "\n".join(
        c.text for t in document.tables for c in t._cells
    )
    assert len(document.inline_shapes) >= 4  # the formulas of the fixed blocks
    # the IP/IK tables of the fixed block are in the draft exactly as in the library
    ip = db.scalars(
        select(TemplateBlock).where(TemplateBlock.key.endswith("influencias_externas"))
    ).one()
    tables = [e["ooxml"] for e in ip.body_template if (e["ooxml"] or "").startswith("<w:tbl")]
    body_tables = {_c14n(t) for t in _root(data).iter(w("tbl"))}
    assert tables and all(_c14n(el) in body_tables for t in tables for el in _fragment(data, t))
    package = zipfile.ZipFile(io.BytesIO(data))
    assert "updateFields" in package.read("word/settings.xml").decode()
    headers = " ".join(package.read(n).decode() for n in package.namelist() if "header" in n)
    assert "EXEMPLO" not in headers.upper()  # the template's technician is not kept
    _assert_package_is_whole(package)


def _assert_package_is_whole(package: zipfile.ZipFile) -> None:
    """What Word checks first: every XML part parses, every relationship resolves."""
    for name in package.namelist():
        if name.endswith((".xml", ".rels")):
            etree.fromstring(package.read(name))
    rels = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    targets = {r.get("Id"): r for r in rels}
    used = set(re.findall(r'r:(?:embed|id|link)="([^"]+)"', package.read(DOCUMENT).decode()))
    assert used <= set(targets), used - set(targets)
    for rid in used:
        rel = targets[rid]
        if rel.get("TargetMode") != "External":
            target = rel.get("Target") or ""
            assert "word/" + target.lstrip("/") in package.namelist(), target


def test_only_writers_assemble(api: Api, db: Session, store: ObjectStore) -> None:
    seed_library(db, store)
    project = api.as_("redator").post("/api/projects", json={"code": "X2", "name": "X"}).json()
    response = api.as_("curador").post(
        f"/api/projects/{project['id']}/documents", json={"type": "MDJ"}
    )
    assert response.status_code == 403


def test_sections_keep_the_block_version(mdj: dict[str, Any], db: Session) -> None:
    rows = db.scalars(select(Section)).all()
    assert all(r.block_version == 1 and r.block_status == "proposed" for r in rows)
