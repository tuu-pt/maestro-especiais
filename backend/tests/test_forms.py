"""Pre-filled forms (Phase 4, task 6): FE loopback, macros kept, no date nor signature, profile."""

import io
import re
import uuid
import zipfile
from typing import Any

import docx
import openpyxl
import pytest
from conftest import Api
from reference_projects import FIXTURES, have_fixtures, load_confirmed
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.assembly.assemble import confirmed_revision
from app.config import Settings
from app.forms import derive
from app.forms.fill import FORMS, template, unique_cells
from app.ingest import ficha_eletrotecnica
from app.ingest.ficha_eletrotecnica import _clean
from app.models import FichaValue, TechnicianProfile
from app.profiles import DEV_PROFILE, personal_terms, profile_for, save_profile

pytestmark = pytest.mark.skipif(not have_fixtures(), reason="sem fixtures de R1")
KEY = "5Q1Kq2y8m0bJb3m4mJ0a0ZKq3f7hC9c7vQnJY2x0Xr8="  # test key (Fernet: 32 bytes, base64)
R1_FE = FIXTURES / "R1/MBERAL/2-PE/Editavel/MBERAL_FichaEletrotécnica_PE_ELE.xlsm"
DATE = re.compile(r"\b\d{1,4}/\d{1,2}/\d{2,4}\b")


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


@pytest.fixture
def r1(api: Api, inline_ingestion: None) -> str:
    return load_confirmed(api, "R1")


def form(api: Api, project_id: str, kind: str) -> bytes:
    response = api.as_("redator").get(f"/api/projects/{project_id}/forms/{kind}")
    assert response.status_code == 200, response.text
    return bytes(response.content)


def docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    return "\n".join(c.text for t in document.tables for r in t.rows for c in unique_cells(r))


# ---------------------------------------------------------------- templates


def test_the_templates_are_the_r1_forms_without_its_values() -> None:
    fe = ficha_eletrotecnica.read(R1_FE.read_bytes())
    derived = derive.derive()
    for kind, data in derived.items():
        # reproducible: make form-templates gives the templates in the repository
        assert _parts(data).keys() == _parts(template(FORMS[kind].template)).keys()
    text = docx_text(derived["identificacao"]) + docx_text(derived["termo"])
    # no value of R1 (ficha, technician, date) is left in the forms
    for c in fe.values:
        if isinstance(c.value, str) and len(c.value) > 3:
            assert c.value not in text, c.key
    assert "@" not in text and not DATE.search(text)
    assert not ficha_eletrotecnica.read(derived["ficha_eletrotecnica"]).values
    sheet = _parts(derived["ficha_eletrotecnica"])
    assert b"mailto:" not in sheet["xl/worksheets/_rels/sheet1.xml.rels"]
    assert b"@" not in sheet["xl/sharedStrings.xml"]


# ---------------------------------------------------------------- ficha eletrotécnica


def test_the_fe_reads_back_as_the_ficha_base(api: Api, r1: str, db: Session) -> None:
    data = form(api, r1, "ficha_eletrotecnica")
    read = {c.key: c.value for c in ficha_eletrotecnica.read(data).values}

    mapped = set(ficha_eletrotecnica.cell_maps()["FE_v.20190222"].cells.values())
    revision = confirmed_revision(db, uuid.UUID(r1))
    assert revision is not None
    revision_values = {
        v.key: _clean(v.key, v.value)
        for v in db.scalars(select(FichaValue).where(FichaValue.revision_id == revision.id))
        if v.key in mapped and v.value not in (None, "", [])
    }
    assert revision_values  # R1 has them
    for key, value in revision_values.items():
        assert read.get(key) == value, key
    assert set(read) <= set(revision_values)


def test_the_fe_keeps_its_macros_and_leaves_the_date_to_the_technician(api: Api, r1: str) -> None:
    data = form(api, r1, "ficha_eletrotecnica")
    parts, source = _parts(data), _parts(R1_FE.read_bytes())

    assert parts["xl/vbaProject.bin"] == source["xl/vbaProject.bin"]
    assert parts["xl/styles.xml"] == source["xl/styles.xml"]
    assert b"extLst" in parts["xl/worksheets/sheet1.xml"]  # the DGEG lists (x14 validations)
    sheet = openpyxl.load_workbook(io.BytesIO(data), keep_vba=True)["Ficha Eletrotecnica"]
    assert sheet["M40"].value is None  # P8: the date is the technician's
    assert sheet["M44"].value == "(Data e assinatura do técnico responsável)"
    assert str(sheet["R29"].value).startswith("=IF(")  # the formulas stay
    # the technician's block: the development profile (R1 was confirmed by dev:tecnico)
    assert sheet["C11"].value == DEV_PROFILE["tec.nome"]
    assert sheet["J12"].value == DEV_PROFILE["tec.email"]


def test_forms_need_a_confirmed_ficha_base(api: Api) -> None:
    project = (
        api.as_("redator").post("/api/projects", json={"code": "X3", "name": "Sem ficha"}).json()
    )
    response = api.as_("redator").get(f"/api/projects/{project['id']}/forms")
    assert response.status_code == 409


# ---------------------------------------------------------------- Identificação and Termo


def test_identificacao_and_termo_are_filled_without_date_or_signature(api: Api, r1: str) -> None:
    listing = {f["kind"]: f for f in api.as_("redator").get(f"/api/projects/{r1}/forms").json()}
    assert set(listing) == {"ficha_eletrotecnica", "identificacao", "termo"}

    for kind in ("identificacao", "termo"):
        text = docx_text(form(api, r1, kind))
        assert "Santo António dos Olivais" in text and "Coimbra" in text  # the ficha-base
        assert DEV_PROFILE["tec.nome"] in text and DEV_PROFILE["tec.oet"] in text  # profile
        assert not DATE.search(text)  # P8
        assert "{{" not in text
        assert "Data e assinatura do técnico responsável" in " ".join(listing[kind]["by_hand"])

    identificacao = docx.Document(io.BytesIO(form(api, r1, "identificacao")))
    rows = {unique_cells(r)[0].text.strip(): [c.text for c in unique_cells(r)]
            for r in identificacao.tables[3].rows}  # fmt: skip
    assert "34,5" in rows["Tensão da RESP [kV]:"]  # the row also has the power to supply
    termo = docx.Document(io.BytesIO(form(api, r1, "termo")))
    nova = next(unique_cells(r) for r in termo.tables[4].rows if "NIP:" in r.cells[0].text)
    assert [c.text for c in nova][-2:] == ["Instalação nova", "X"]  # ele.instalacao = Nova


def test_downloading_a_form_is_audited(api: Api, r1: str) -> None:
    form(api, r1, "termo")
    events = api.as_("redator").get(f"/api/projects/{r1}/audit").json()
    assert "Descarregou o Termo de Responsabilidade pré-preenchido" in [
        e["description"] for e in events
    ]


# ---------------------------------------------------------------- profile


def test_the_profile_is_encrypted_at_rest(db: Session) -> None:
    settings = Settings(profile_encryption_key=KEY)
    save_profile(db, settings, "tec:1", {"tec.nome": "Nome Muito Particular", "tec.oet": "12345"})

    row = db.scalars(select(TechnicianProfile).where(TechnicianProfile.user_id == "tec:1")).one()
    assert "Particular" not in row.data and "12345" not in row.data
    assert profile_for(db, settings, "tec:1")["tec.nome"] == "Nome Muito Particular"
    assert ("profile", "Nome Muito Particular") in personal_terms(db, settings)
    assert profile_for(db, Settings(), "tec:1") == {}  # without the key, nothing


def test_the_technician_edits_the_profile_and_personal_fields_stay_masked(
    api: Api, settings: Settings
) -> None:
    settings.profile_encryption_key = KEY
    before = api.as_("tecnico").get("/api/me/profile").json()
    assert before["source"] == "development"

    saved = api.as_("tecnico").put(
        "/api/me/profile",
        json={"values": {"tec.nome": "Outro Nome", "doc.local": "Aveiro"}},
    )

    assert saved.status_code == 200
    fields: dict[str, Any] = {f["key"]: f for f in saved.json()["fields"]}
    assert saved.json()["source"] == "saved"
    assert fields["tec.nome"] == {**fields["tec.nome"], "filled": True, "value": None}
    assert fields["doc.local"]["value"] == "Aveiro"
    assert "Outro Nome" not in saved.text
    assert api.as_("redator").get("/api/me/profile").status_code == 403
    bad = api.as_("tecnico").put("/api/me/profile", json={"values": {"id.requerente.nome": "x"}})
    assert bad.status_code == 422
