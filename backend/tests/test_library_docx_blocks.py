"""MDJ/CTE split into sections with the original OOXML, and the round trip (Phase 3, task 2)."""

import io
import zipfile
from pathlib import Path

import docx
import pytest
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.library import docx_blocks
from app.library.sources import media_key, package_key, reference_documents, seed_sources
from app.models import SourceDocument
from app.storage import ObjectStore

FIXTURES = Path(__file__).resolve().parents[2] / "data" / "fixtures"


def documents() -> list[tuple[str, str, str, bytes]]:
    if not (FIXTURES / "R1").is_dir() or not (FIXTURES / "R2").is_dir():
        return []
    return [
        (code, doc_type, file, data)
        for code in ("R1", "R2")
        for doc_type, file, data in reference_documents(FIXTURES, code)
    ]


DOCS = documents()
pytestmark = pytest.mark.skipif(not DOCS, reason="data/fixtures/R1 e R2 são precisos")


def one(code: str, doc_type: str) -> docx_blocks.SplitDocument:
    [data] = [d for c, t, _, d in DOCS if (c, t) == (code, doc_type)]
    return docx_blocks.split(data)


def titles(parts: docx_blocks.SplitDocument, level: int) -> list[str]:
    return [s.title for s in parts.sections if s.kind == "block" and s.level == level]


def test_each_reference_project_has_one_mdj_and_one_cte() -> None:
    assert sorted((c, t) for c, t, _, _ in DOCS) == [
        ("R1", "CTE"), ("R1", "MDJ"), ("R2", "CTE"), ("R2", "MDJ"),
    ]  # fmt: skip


def test_mdj_skeleton_of_r1() -> None:
    parts = one("R1", "MDJ")

    assert [s.kind for s in parts.sections[:2]] == ["cover", "index"]
    assert parts.sections[-1].kind == "signature"
    assert titles(parts, 1) == [
        "INTRODUÇÃO", "LEGISLAÇÃO E NORMAS",
        "CARACTERÍSTICAS DOS EQUIPAMENTOS EM FUNÇÃO DAS INFLUÊNCIAS EXTERNAS",
        "CLASSIFICAÇÃO QUANTO À UTILIZAÇÃO DO LOCAL",
        "INSTALAÇÃO DE ALIMENTAÇÃO, DISTRIBUIÇÃO E MEDIDA DE ENERGIA", "DIMENSIONAMENTO ELÉTRICO",
        "QUADRO ELÉTRICO", "CANALIZAÇÕES", "CAIXAS", "INSTALAÇÕES ELÉTRICAS A CONSIDERAR",
        "PROTEÇÃO DOS UTILIZADORES", "DÚVIDAS E CASOS OMISSOS",
    ]  # fmt: skip
    assert len(titles(parts, 2)) == 19
    keys = {s.key for s in parts.sections}
    assert "canalizacoes.canalizacoes_enterradas" in keys  # only in R1
    assert "dimensionamento_eletrico.quedas_de_tensao" in keys
    assert parts.warnings == []


def test_mdj_of_r2_has_the_blocks_only_r2_has() -> None:
    keys = {s.key for s in one("R2", "MDJ").sections}

    assert {
        "regulamento_dos_produtos_de_construcao_rpc",
        "sistema_automatico_de_detecao_de_incendio_sadi.matriz_de_incendio",
        "instalacao_fotovoltaica", "carregamento_de_veiculos_eletricos",
        "instalacao_audiovisual_auditorio",
    } <= keys  # fmt: skip
    assert "canalizacoes.canalizacoes_enterradas" not in keys


def test_cover_index_and_signature() -> None:
    parts = one("R1", "MDJ")
    cover, index = parts.sections[0], parts.sections[1]
    signature = parts.sections[-1]

    assert "MEMÓRIA DESCRITIVA E JUSTIFICATIVA" in cover.text
    assert index.text.startswith("Conteúdo")
    assert signature.text.splitlines()[:2] == ["Coimbra, junho de 2026", "O Técnico, "]
    # the text before the signature stays in "dúvidas e casos omissos"
    assert "casos omissos" in parts.sections[-2].text


def test_an_empty_one_by_two_table_is_not_a_level_1_band() -> None:
    parts = one("R1", "CTE")

    assert titles(parts, 1) == [
        "CONDIÇÕES TÉCNICAS GERAIS", "CONDIÇÕES TÉCNICAS ESPECIAIS", "DÚVIDAS E CASOS OMISSOS",
    ]  # fmt: skip
    cables = next(s for s in parts.sections if s.key.endswith(".cabos_e_fios"))
    assert cables.stats()["tables"] == 1  # the empty band is content of "Cabos e Fios"


def test_ip_ik_tables_stay_in_their_section() -> None:
    for code in ("R1", "R2"):
        section = next(
            s for s in one(code, "MDJ").sections
            if s.key == "caracteristicas_dos_equipamentos_em_funcao_das_influencias_externas"
        )  # fmt: skip
        assert section.stats()["tables"] == 3  # the band and the two tables of codes
        assert "Classe de influências externas" in section.text


@pytest.mark.parametrize(("code", "doc_type"), [(c, t) for c, t, _, _ in DOCS])
def test_round_trip_gives_the_same_document(code: str, doc_type: str) -> None:
    [data] = [d for c, t, _, d in DOCS if (c, t) == (code, doc_type)]
    parts = docx_blocks.split(data)

    rebuilt = docx_blocks.rebuild(parts.package, [s.ooxml for s in parts.sections])

    assert docx_blocks.canonical_document(rebuilt) == docx_blocks.canonical_document(data)
    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(rebuilt))
    assert after.namelist() == before.namelist()
    for name in before.namelist():
        if name != docx_blocks.DOCUMENT:
            assert after.read(name) == before.read(name), name
    original, again = docx.Document(io.BytesIO(data)), docx.Document(io.BytesIO(rebuilt))
    assert [p.text for p in again.paragraphs] == [p.text for p in original.paragraphs]
    assert len(again.tables) == len(original.tables)
    assert len(again.inline_shapes) == len(original.inline_shapes)
    styles = {p.style.style_id for p in again.paragraphs if p.style is not None}
    assert styles == {p.style.style_id for p in original.paragraphs if p.style is not None}


@pytest.mark.parametrize(("code", "doc_type"), [(c, t) for c, t, _, _ in DOCS])
def test_every_relationship_a_section_uses_is_known(code: str, doc_type: str) -> None:
    parts = one(code, doc_type)

    images = 0
    for section in parts.sections:
        for rel in section.rels.values():
            if rel.type == "image":
                images += 1
                assert rel.sha256 in parts.media
    assert images == sum(len(s.rels) for s in parts.sections if s.rels)  # only images, here
    assert parts.warnings == []


def test_sections_add_up_to_the_whole_body() -> None:
    for _, _, _, data in DOCS:
        parts = docx_blocks.split(data)
        body = docx.Document(io.BytesIO(data)).element.body
        assert sum(s.stats()["elements"] for s in parts.sections) == len(body) - 1  # sectPr
        assert [s.order for s in parts.sections] == list(range(1, len(parts.sections) + 1))


def test_not_a_word_document() -> None:
    with pytest.raises(docx_blocks.DocxError):
        docx_blocks.split(b"not a zip")


# ---------------------------------------------------------------- stored in the library


def test_seed_stores_packages_media_and_sections(db: Session, store: ObjectStore) -> None:
    summary = seed_sources(db, store, FIXTURES, ("R1", "R2"))

    assert summary["source_documents"] == 4
    rows = db.scalars(select(SourceDocument).order_by(SourceDocument.project_code)).all()
    assert {(d.project_code, d.doc_type) for d in rows} == {
        ("R1", "MDJ"), ("R1", "CTE"), ("R2", "MDJ"), ("R2", "CTE"),
    }  # fmt: skip
    mdj = next(d for d in rows if (d.project_code, d.doc_type) == ("R1", "MDJ"))
    assert mdj.package_key == package_key(mdj.sha256) and "MBERAL" not in mdj.package_key
    rebuilt = docx_blocks.rebuild(store.get(mdj.package_key), [s.ooxml for s in mdj.sections])
    [data] = [d for c, t, _, d in DOCS if (c, t) == ("R1", "MDJ")]
    assert docx_blocks.canonical_document(rebuilt) == docx_blocks.canonical_document(data)
    image = next(r for s in mdj.sections for r in s.rels.values() if r["type"] == "image")
    assert store.get(media_key(image["sha256"]))
    assert mdj.sections[-1].kind == "signature"


def test_seeding_again_changes_nothing(db: Session, store: ObjectStore) -> None:
    seed_sources(db, store, FIXTURES, ("R1",))
    ids = {d.id for d in db.scalars(select(SourceDocument))}

    again = seed_sources(db, store, FIXTURES, ("R1",))

    assert {d.id for d in db.scalars(select(SourceDocument))} == ids
    assert again["source_documents"] == 2
