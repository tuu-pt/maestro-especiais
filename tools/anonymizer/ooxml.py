"""docx / xlsx / xlsm anonymization at ZIP/XML level.

The package is never re-saved through python-docx or openpyxl: parts are edited in
place, so VBA, formatting and the cached values of formulas survive. Text split across
runs is handled paragraph by paragraph; every other text node and attribute is also
processed (headers, comments, docProps, sheet cells, relationships...).
"""

import io
import re
import zipfile
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

import docx
import openpyxl
from lxml import etree

from anonymizer import images
from anonymizer.engine import Seed, TextAnonymizer
from anonymizer.findings import Finding
from anonymizer.harvest import dedupe, fe_maps, seeds_for, seeds_from_grid

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

EXTENSIONS = (".docx", ".docm", ".dotx", ".xlsx", ".xlsm", ".xltx")
_PACKAGE_EXTS = (*EXTENSIONS, ".pptx")
_XML_EXTS = (".xml", ".rels", ".vml")
_IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf", ".svg")
_THUMBNAIL = "docprops/thumbnail"
_QUIET_BINARIES = ("vbaproject.bin", "printersettings", ".odttf", ".ttf", ".fntdata")

# (container, label, text tags grouped per container)
_GROUPS = (
    (f"{{{W}}}p", "parágrafo", (f"{{{W}}}t", f"{{{W}}}delText", f"{{{W}}}instrText")),
    (f"{{{S}}}si", "texto partilhado", (f"{{{S}}}t",)),
    (f"{{{S}}}is", "texto", (f"{{{S}}}t",)),
    (f"{{{A}}}p", "parágrafo", (f"{{{A}}}t",)),
)

Edit = tuple[int, int, str]
TextFn = Callable[[str, str], list[Edit]]

_PARSER = etree.XMLParser(
    resolve_entities=False, huge_tree=True, remove_blank_text=False, no_network=True
)


# ---------------------------------------------------------------- text editing


def splice(segments: list[str], edits: list[Edit]) -> list[str]:
    """Apply edits on the concatenation of segments, keeping the segment structure.

    A replacement goes into the first segment it touches; the rest of the span is
    removed from the following segments.
    """
    starts, pos = [], 0
    for seg in segments:
        starts.append(pos)
        pos += len(seg)
    out = list(segments)
    for s, e, rep in sorted(edits, reverse=True):
        first = True
        for j, seg in enumerate(segments):
            seg_start, seg_end = starts[j], starts[j] + len(seg)
            if seg_end <= s or seg_start >= e or seg_start == seg_end:
                continue
            ls, le = max(s, seg_start) - seg_start, min(e, seg_end) - seg_start
            out[j] = out[j][:ls] + (rep if first else "") + out[j][le:]
            first = False
    return out


def _apply(text: str, edits: list[Edit]) -> str:
    return splice([text], edits)[0]


def _owner(node: Any, tag: str) -> Any:
    parent = node.getparent()
    while parent is not None and parent.tag != tag:
        parent = parent.getparent()
    return parent


def _cell_ref(el: Any) -> str | None:
    while el is not None:
        if el.tag == f"{{{S}}}c":
            ref = el.get("r")
            return str(ref) if ref else None
        el = el.getparent()
    return None


def _local(tag: Any) -> str:
    return etree.QName(tag).localname if isinstance(tag, str) else "comentário"


# Excel header/footer codes (&C, &"Arial,Bold", &12...) are glued to the text around them.
_HEADER_FOOTER = {
    "oddHeader",
    "oddFooter",
    "evenHeader",
    "evenFooter",
    "firstHeader",
    "firstFooter",
}
_HF_CODE = re.compile(r'&"[^"]*"|&\d+|&[A-Za-z&]')


def _segmented(fn: TextFn, text: str, where: str) -> list[Edit]:
    """Run fn on the pieces of text between header/footer codes."""
    edits: list[Edit] = []
    pos = 0
    for m in [*_HF_CODE.finditer(text), None]:
        end = m.start() if m else len(text)
        if text[pos:end].strip():
            edits += [(pos + a, pos + b, rep) for a, b, rep in fn(text[pos:end], where)]
        pos = m.end() if m else end
    return edits


def walk_xml(root: Any, member: str, fn: TextFn) -> bool:
    """Run fn over every text of an XML part and apply its edits. Returns True if changed."""
    changed = False
    handled: set[Any] = set()
    for container_tag, label, text_tags in _GROUPS:
        for i, el in enumerate(root.iter(container_tag), 1):
            ref = _cell_ref(el)
            where = f"{member} · célula {ref}" if ref else f"{member} · {label} {i}"
            for tag in text_tags:
                nodes = [n for n in el.iter(tag) if _owner(n, container_tag) is el]
                if not nodes:
                    continue
                handled.update(nodes)
                texts = [n.text or "" for n in nodes]
                edits = fn("".join(texts), where)
                if not edits:
                    continue
                for node, new in zip(nodes, splice(texts, edits), strict=True):
                    if new != (node.text or ""):
                        node.text = new
                        if new != new.strip():
                            node.set(XML_SPACE, "preserve")
                changed = True
    for el in root.iter():
        if isinstance(el, etree._ProcessingInstruction):
            continue
        ref = _cell_ref(el)
        where = f"{member} · célula {ref}" if ref else f"{member} · {_local(el.tag)}"
        if el not in handled and el.text and el.text.strip():
            is_hf = isinstance(el.tag, str) and _local(el.tag) in _HEADER_FOOTER
            edits = _segmented(fn, el.text, where) if is_hf else fn(el.text, where)
            if edits:
                el.text = _apply(el.text, edits)
                changed = True
        if el.tail and el.tail.strip():
            edits = fn(el.tail, where)
            if edits:
                el.tail = _apply(el.tail, edits)
                changed = True
        if isinstance(el.tag, str):
            for name, value in el.attrib.items():
                if value.strip():
                    edits = fn(value, f"{where} · atributo {_local(name)}")
                    if edits:
                        el.set(name, _apply(value, edits))
                        changed = True
    return changed


def _drop_thumbnail_relationship(raw: bytes) -> bytes:
    root = etree.fromstring(raw, _PARSER)
    dropped = False
    for rel in list(root):
        if _THUMBNAIL in str(rel.get("Target", "")).lower():
            root.remove(rel)
            dropped = True
    return _serialize(root) if dropped else raw


def _serialize(root: Any) -> bytes:
    tree = root.getroottree()
    return bytes(
        etree.tostring(
            tree, xml_declaration=True, encoding="UTF-8", standalone=tree.docinfo.standalone
        )
    )


# ---------------------------------------------------------------- package traversal


def _members(data: bytes) -> Iterator[tuple[zipfile.ZipInfo, bytes]]:
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        for info in package.infolist():
            yield info, package.read(info)


def _is_image(name: str) -> bool:
    return name.endswith(_IMAGE_EXTS) or "/media/" in name


def _binary_findings(data: bytes, where: str, name: str, engine: TextAnonymizer) -> list[Finding]:
    hits = sorted(set(engine.binary_hits(data)))
    if hits:
        return [Finding("pii_in_binary", where, kind) for kind in hits]
    if any(q in name for q in _QUIET_BINARIES):
        return []
    if "embeddings/" in name or name.endswith(".bin"):
        return [Finding("binary_not_inspected", where)]
    return []


def transform_package(
    data: bytes, engine: TextAnonymizer, strip_images: bool, prefix: str = ""
) -> tuple[bytes, list[Finding]]:
    findings: list[Finding] = []

    def fn(text: str, where: str) -> list[Edit]:
        _, reps = engine.anonymize(text)
        for r in reps:
            findings.append(
                Finding("replaced", where, r.kind, original=r.original, pseudonym=r.pseudonym)
            )
        return [(r.start, r.end, r.pseudonym) for r in reps]

    counts = {"images": 0, "blanked": 0, "metadata": 0}
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as out:
        for info, raw in _members(data):
            name = info.filename.lower()
            where = prefix + info.filename
            new = raw
            if name.startswith(_THUMBNAIL):
                # A rendering of the first page: dropped, with its relationship.
                findings.append(Finding("thumbnail_removed", where))
                continue
            if name == "_rels/.rels":
                raw = new = _drop_thumbnail_relationship(raw)
            if name.endswith(_XML_EXTS):
                try:
                    root = etree.fromstring(raw, _PARSER)
                except etree.XMLSyntaxError:
                    findings.append(Finding("unreadable_part", where))
                else:
                    if walk_xml(root, where, fn):
                        new = _serialize(root)
            elif name.endswith(_PACKAGE_EXTS):
                new, inner = transform_package(raw, engine, strip_images, f"{where} › ")
                findings += inner
            elif _is_image(name):
                counts["images"] += 1
                if strip_images:
                    blank = images.blank(raw)
                    if blank is None:
                        findings.append(Finding("image_not_blanked", where))
                    else:
                        new = blank
                        counts["blanked"] += 1
                else:
                    new = images.strip_metadata(raw)
                    counts["metadata"] += new != raw
            else:
                findings += _binary_findings(raw, where, name, engine)
            out.writestr(info, new)
    if counts["blanked"]:
        findings.append(Finding("images_blanked", prefix, count=counts["blanked"]))
    if counts["metadata"]:
        findings.append(Finding("image_metadata_removed", prefix, count=counts["metadata"]))
    remaining = counts["images"] - counts["blanked"]
    if remaining and not strip_images:
        findings.append(Finding("images_present", prefix, count=remaining))
    return buffer.getvalue(), findings


def visit_package(
    data: bytes, fn: TextFn, engine: TextAnonymizer | None = None, prefix: str = ""
) -> list[Finding]:
    """Read-only traversal: fn sees every text; binaries are checked when engine is given."""
    findings: list[Finding] = []
    for info, raw in _members(data):
        name = info.filename.lower()
        where = prefix + info.filename
        if name.endswith(_XML_EXTS):
            try:
                root = etree.fromstring(raw, _PARSER)
            except etree.XMLSyntaxError:
                findings.append(Finding("unreadable_part", where))
                continue
            walk_xml(root, where, fn)
        elif name.endswith(_PACKAGE_EXTS):
            findings += visit_package(raw, fn, engine, f"{where} › ")
        elif engine is not None and not _is_image(name):
            findings += [
                f for f in _binary_findings(raw, where, name, engine) if f.code == "pii_in_binary"
            ]
    return findings


# ---------------------------------------------------------------- file-level API


def transform_file(
    src: Path, dst: Path, engine: TextAnonymizer, strip_images: bool
) -> list[Finding]:
    data, findings = transform_package(src.read_bytes(), engine, strip_images)
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_bytes(data)
    return findings


def texts(path: Path) -> list[tuple[str, str]]:
    collected: list[tuple[str, str]] = []

    def fn(text: str, where: str) -> list[Edit]:
        collected.append((where, text))
        return []

    visit_package(path.read_bytes(), fn)
    return collected


def scan_file(path: Path, checker: TextAnonymizer) -> list[Finding]:
    findings: list[Finding] = []

    def fn(text: str, where: str) -> list[Edit]:
        for r in checker.residuals(text):
            code = "residual" if r.severity == "error" else "possible_name"
            findings.append(Finding(code, where, r.kind, original=r.value))
        return []

    findings += visit_package(path.read_bytes(), fn, checker)
    return findings


# ---------------------------------------------------------------- harvest


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _docx_grids(path: Path) -> list[list[list[str]]]:
    grids: list[list[list[str]]] = []

    def visit(tables: Any) -> None:
        for table in tables:
            grids.append([[cell.text for cell in row.cells] for row in table.rows])
            for row in table.rows:
                for cell in row.cells:
                    visit(cell.tables)

    visit(docx.Document(str(path)).tables)
    return grids


def harvest(path: Path) -> tuple[list[Seed], list[Finding], bool]:
    """Seeds from form tables and, for the ficha eletrotécnica, its fixed cells.

    Returns (seeds, findings, is_form). A file is a form when its tables hold at least
    two labelled personal fields.
    """
    findings: list[Finding] = []
    seeds: list[Seed] = []
    grid_seeds: list[Seed] = []
    suffix = path.suffix.lower()
    if suffix in {".docx", ".dotx"}:
        for grid in _docx_grids(path):
            grid_seeds += seeds_from_grid(grid)
    elif suffix in {".xlsx", ".xlsm", ".xltx"}:
        workbook = openpyxl.load_workbook(path, data_only=True, keep_links=False)
        for sheet in workbook.worksheets:
            rows = [[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            grid_seeds += seeds_from_grid(rows)
            seeds += _fe_seeds(sheet, findings)
        workbook.close()
    return dedupe(seeds + grid_seeds), findings, len(grid_seeds) >= 2


def _fe_seeds(sheet: Any, findings: list[Finding]) -> list[Seed]:
    maps = fe_maps()
    for fe in maps:
        version = _cell_text(sheet[fe.version_cell].value).strip()
        if version == fe.template_version:
            seeds: list[Seed] = []
            for cell, kind in fe.cells.items():
                seeds += seeds_for(kind, _cell_text(sheet[cell].value))
            return seeds
    cells = {m.version_cell for m in maps}
    for cell in cells:
        version = _cell_text(sheet[cell].value).strip()
        if version.startswith("FE_v"):
            findings.append(Finding("unknown_fe_version", f"{sheet.title}!{cell}"))
    return []
