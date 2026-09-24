"""Synthetic reference project that imitates the structure of R1/R2 (SPEC 8.2, 12.2).

All personal data here is invented. Files are built at test time, never committed.
"""

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import docx
import openpyxl
import pymupdf
from docx.document import Document as DocxDocument
from docx.oxml import parse_xml
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from openpyxl.comments import Comment

from anonymizer.detectors import nif_check_digit
from anonymizer.textnorm import digits_only, fold


def make_nif(first8: str) -> str:
    return first8 + str(nif_check_digit(first8))


@dataclass(frozen=True)
class Person:
    name: str
    nif: str
    address: str
    postal_code: str
    phone: str
    email: str


PROMOTOR = Person(
    name="João Pedro Almeida",
    nif=make_nif("23456789"),
    address="Rua das Flores, n.º 12, 3.º Esq.",
    postal_code="4000-123",
    phone="912 345 678",
    email="joao.almeida@mail.pt",
)
TECNICO = Person(
    name="Maria Sousa Ferreira",
    nif=make_nif("18765432"),
    address="Avenida da Boavista 1234",
    postal_code="4100-456",
    phone="+351 225 123 456",
    email="maria.ferreira@eng.pt",
)
OTHER_REQUERENTE = Person(  # C7: the ficha of the second project belongs to someone else
    name="Carlos Alberto Mendes",
    nif=make_nif("29876543"),
    address="Travessa de São Victor 7",
    postal_code="4000-987",
    phone="933 222 111",
    email="carlos.mendes@mail.pt",
)
TECNICO_CC = "12345678 9 ZX4"
TECNICO_DGEG = "4321"
TECNICO_OET_FORMS = "15678"
TECNICO_OET_MDJ = "15687"  # C3: differs between the MDJ and the forms
GPS = "41.1579, -8.6291"
DESENHADOR = "Rui Manuel Costa"

NAMES = [PROMOTOR.name, TECNICO.name, OTHER_REQUERENTE.name, DESENHADOR]
LONG_NUMBERS = (
    [p.nif for p in (PROMOTOR, TECNICO, OTHER_REQUERENTE)]
    + [p.phone for p in (PROMOTOR, TECNICO, OTHER_REQUERENTE)]
    + [TECNICO_CC]
)
SHORT_NUMBERS = [TECNICO_DGEG, TECNICO_OET_FORMS, TECNICO_OET_MDJ]
LITERALS = [
    *(p.email for p in (PROMOTOR, TECNICO, OTHER_REQUERENTE)),
    *(p.postal_code for p in (PROMOTOR, TECNICO, OTHER_REQUERENTE)),
    "Rua das Flores",
    "Avenida da Boavista",
    "Travessa de São Victor",
    "41.1579",
    "8.6291",
]
# Must survive anonymization untouched (not personal data).
CONTROLS = [
    "34,5 kVA",
    "H07V-U",
    "XZ1(frt,zh)",
    "Portaria n.º 949-A/2006",
    "DL 96/2017",
    "IP65",
    "16A-250V",
    "Cedofeita",
    "Porto",
    "Moradia unifamiliar",
]


def spaced(nif: str) -> str:
    return f"{nif[:3]} {nif[3:6]} {nif[6:]}"


# ---------------------------------------------------------------- absence check


def name_variants(name: str) -> list[str]:
    tokens = name.split()
    return [name, name.upper(), f"{tokens[0]} {tokens[-1]}", f"{tokens[0][0]}. {tokens[-1]}"]


def find_pii(text: str) -> list[str]:
    """Which synthetic personal values occur in text (accent/case-insensitive)."""
    folded = fold(text)[0]
    hits = []
    for name in NAMES:
        for variant in name_variants(name):
            if re.search(rf"(?<![0-9a-z]){re.escape(fold(variant)[0])}(?![0-9a-z])", folded):
                hits.append(variant)
    for number in LONG_NUMBERS:
        digits = digits_only(number)
        digits = digits[:9] if number == TECNICO_CC else digits[-9:]
        pattern = r"(?<!\d)" + r"[ .\u00a0-]?".join(digits) + r"(?!\d)"
        if re.search(pattern, text):
            hits.append(number)
    for number in SHORT_NUMBERS:
        if re.search(rf"(?<![\d.,]){number}(?![\d]|[.,]\d)", text):
            hits.append(number)
    for literal in LITERALS:
        if fold(literal)[0] in folded:
            hits.append(literal)
    return hits


# ---------------------------------------------------------------- images


def png(color: tuple[int, int, int] = (20, 20, 160), text_chunk: bytes = b"") -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    pix.set_rect(pix.irect, color)
    data = bytes(pix.tobytes("png"))
    if not text_chunk:
        return data
    import struct
    import zlib

    body = b"tEXt" + b"Author\x00" + text_chunk
    chunk = struct.pack(">I", len(body) - 4) + body + struct.pack(">I", zlib.crc32(body))
    return data[:33] + chunk + data[33:]  # right after IHDR


def jpeg_with_exif(text: bytes) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    pix.set_rect(pix.irect, (200, 30, 30))
    data = bytes(pix.tobytes("jpeg"))
    if not text:
        return data
    payload = b"Exif\x00\x00" + text
    segment = b"\xff\xe1" + (len(payload) + 2).to_bytes(2, "big") + payload
    return data[:2] + segment + data[2:]


# ---------------------------------------------------------------- zip helpers


def add_members(path: Path, members: dict[str, bytes], content_types: str = "") -> None:
    """Add raw members to an OOXML package (and optional <Default>/<Override> entries)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "[Content_Types].xml" and content_types:
                data = data.replace(b"</Types>", content_types.encode() + b"</Types>")
            out.writestr(info, data)
        for name, data in members.items():
            out.writestr(name, data)
    path.write_bytes(buffer.getvalue())


def replace_in_member(path: Path, member: str, old: bytes, new: bytes) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        assert member in src.namelist(), f"{member} not in package"
        for info in src.infolist():
            data = src.read(info)
            if info.filename == member:
                assert old in data, f"{old!r} not in {member}"
                data = data.replace(old, new)
            out.writestr(info, data)
    path.write_bytes(buffer.getvalue())


# ---------------------------------------------------------------- ficha eletrotécnica


VBA_STUB = b"\xd0\xcf\x11\xe0VBA-PROJECT-STUB" + b"\x00" * 64


def ficha_eletrotecnica(path: Path, requerente: Person = PROMOTOR) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Ficha Eletrotecnica"
    ws["B5"] = "Nome"
    ws["C5"] = requerente.name
    ws["P5"] = "NIF"
    ws["Q5"] = int(requerente.nif)  # stored as a number, as Excel users do
    ws["J6"] = "Moradia unifamiliar"
    ws["C7"] = f"{requerente.address}, {requerente.postal_code} Porto"
    ws["C8"] = requerente.email
    ws["C15"] = "Rua das Flores"
    ws["M15"] = "4000-123"
    ws["Q15"] = "Cedofeita"
    ws["G16"] = "Porto"
    ws["Q16"] = GPS
    ws["B29"] = "C"
    ws["P29"] = 34.5
    ws["Q29"] = 1
    ws["S29"] = "=P29*1.2"
    ws["I40"] = TECNICO.name
    ws["I42"] = TECNICO_DGEG
    ws["I44"] = f"Tel. {TECNICO.phone}"
    ws["R45"] = "FE_v.20190222"
    ws["B30"] = "Observações"
    ws["B30"].comment = Comment(f"Ligar ao técnico: {TECNICO.phone}", TECNICO.name)
    assert ws.oddHeader is not None
    ws.oddHeader.center.text = requerente.name
    wb.properties.creator = TECNICO.name
    wb.properties.lastModifiedBy = TECNICO.name
    wb.save(path)
    replace_in_member(
        path, "xl/worksheets/sheet1.xml", b"<f>P29*1.2</f><v></v>", b"<f>P29*1.2</f><v>41.4</v>"
    )
    add_members(
        path,
        {"xl/vbaProject.bin": VBA_STUB},
        '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>',
    )
    return path


def tabela_calculo(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tabela"
    ws.append(
        [
            "ORIGEM",
            "DESTINO",
            "kVA",
            "IB (A)",
            "In (A)",
            "IΔn (mA)",
            "Iz (A)",
            "I2 (A)",
            "1,45·Iz (A)",
            "CABO",
        ]
    )
    ws.append(["ENTRADA DE ENERGIA"])
    ws.append(
        ["Portinhola", "Q.E.G.", 34.5, 50, 63, None, 347.17, 504, "=G3*1.45", "XZ1(frt,zh) 4x16"]
    )
    ws.append(["Q.E.G.", "Q.P.1", 12, 18, 25, 30, 32, 40, "=G4*1.45", "H07V-U 3G6"])
    ws["A10"] = f"Projetista: {TECNICO.name}"
    wb.save(path)
    replace_in_member(
        path, "xl/worksheets/sheet1.xml", b"<f>G3*1.45</f><v></v>", b"<f>G3*1.45</f><v>503.4</v>"
    )
    return path


# ---------------------------------------------------------------- forms and MDJ


def _add_header_footer(document: DocxDocument, header: str, footer: str) -> None:
    section = document.sections[0]
    section.header.paragraphs[0].text = header
    section.footer.paragraphs[0].text = footer


def _add_textbox(paragraph: Paragraph, text: str) -> None:
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape style="width:200pt;height:40pt">'
        f"<v:textbox><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
        "</v:textbox></v:shape></w:pict></w:r>"
    )
    paragraph._p.append(parse_xml(xml))


def identificacao(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Identificação do projeto", level=1)
    d.add_paragraph("Promotor")
    t = d.add_table(rows=0, cols=2)
    for label, value in [
        ("Nome", PROMOTOR.name),
        ("NIF", spaced(PROMOTOR.nif)),
        ("Morada", PROMOTOR.address),
        ("Código postal", f"{PROMOTOR.postal_code} Porto"),
        ("Telefone", PROMOTOR.phone),
        ("Email", PROMOTOR.email),
    ]:
        row = t.add_row().cells
        row[0].text, row[1].text = label, value
    d.add_paragraph("Técnico responsável")
    t2 = d.add_table(rows=2, cols=4)
    for c, (label, value) in enumerate(
        [
            ("Nome", TECNICO.name),
            ("N.º DGEG", TECNICO_DGEG),
            ("OET n.º", TECNICO_OET_FORMS),
            ("Cartão de Cidadão", TECNICO_CC),
        ]
    ):
        t2.cell(0, c).text, t2.cell(1, c).text = label, value
    p = d.add_paragraph("Requerente: ")
    for part in ("Jo", "ão Pedro ", "Almeida"):  # a name split across three runs
        p.add_run(part)
    run = d.add_paragraph().add_run("Nota interna")
    d.add_comment(run, text=f"Confirmar NIF {TECNICO.nif}", author=TECNICO.name, initials="MF")
    _add_textbox(d.add_paragraph(), f"Tel. {TECNICO.phone}")
    d.add_paragraph("Local e data: ____________  Assinatura: ____________")
    d.add_picture(io.BytesIO(png(text_chunk=TECNICO.name.encode("latin-1"))), width=Cm(3))
    _add_header_footer(d, f"{PROMOTOR.name} – NIF {PROMOTOR.nif}", TECNICO.email)
    d.core_properties.author = TECNICO.name
    d.core_properties.last_modified_by = TECNICO.name
    d.save(str(path))
    return path


def termo(path: Path) -> Path:
    d = docx.Document()
    d.add_heading("Termo de responsabilidade", level=1)
    d.add_paragraph(
        f"{TECNICO.name}, contribuinte n.º {TECNICO.nif}, portadora do Cartão de Cidadão "
        f"n.º {TECNICO_CC}, inscrita na DGEG com o n.º {TECNICO_DGEG} e na Ordem dos "
        f"Engenheiros Técnicos com o n.º {TECNICO_OET_FORMS}, com morada na "
        f"{TECNICO.address}, {TECNICO.postal_code} Porto, telefone {TECNICO.phone}, "
        f"email {TECNICO.email}, declara que o projeto de instalações elétricas da obra "
        f"requerida por {PROMOTOR.name} cumpre as normas aplicáveis."
    )
    t = d.add_table(rows=0, cols=2)
    for label, value in [("Nome", TECNICO.name), ("NIF", TECNICO.nif), ("Email", TECNICO.email)]:
        row = t.add_row().cells
        row[0].text, row[1].text = label, value
    d.add_paragraph("Data: ____________  Assinatura: ____________")
    d.core_properties.author = TECNICO.name
    d.save(str(path))
    return path


def mdj(path: Path, *, embedded_xlsx: bool = True) -> Path:
    d = docx.Document()
    d.add_heading("Memória descritiva e justificativa", level=1)
    d.add_paragraph(
        "A instalação destina-se a uma Moradia unifamiliar com potência a alimentar de "
        "34,5 kVA, alimentada a partir da portinhola, com cabos XZ1(frt,zh) e fios H07V-U."
    )
    d.add_paragraph(
        "São aplicáveis a Portaria n.º 949-A/2006 e o DL 96/2017. As tomadas são de 16A-250V "
        "com grau de proteção IP65. Freguesia de Cedofeita, concelho do Porto."
    )
    d.add_paragraph(f"O requerente, {PROMOTOR.name.upper()}, reside na {PROMOTOR.address}.")
    d.add_paragraph(f"Coordenadas GPS: {GPS}.")
    d.add_paragraph(
        f"O técnico, Eng.ª {TECNICO.name.split()[0]} {TECNICO.name.split()[-1]}, "
        f"OET n.º {TECNICO_OET_MDJ}. Revisto por M. Ferreira."
    )
    plain = jpeg_with_exif(b"")
    d.add_picture(io.BytesIO(plain), width=Cm(3))
    d.core_properties.author = TECNICO.name
    d.save(str(path))
    # python-docx refuses to parse a fake EXIF block, so the photo gets it afterwards.
    with zipfile.ZipFile(path) as package:
        media = next(n for n in package.namelist() if n.startswith("word/media/"))
    replace_in_member(
        path, media, plain, jpeg_with_exif(f"GPS {GPS} {TECNICO.name}".encode("latin-1"))
    )
    if embedded_xlsx:
        inner = path.with_suffix(".inner.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws["A1"] = f"Requerente: {PROMOTOR.name}"
        wb.save(inner)
        add_members(path, {"word/embeddings/Microsoft_Excel_Worksheet.xlsx": inner.read_bytes()})
        inner.unlink()
    return path


def add_ole_with_pii(path: Path) -> None:
    """An embedded OLE object the anonymizer cannot edit, holding a NIF in UTF-16."""
    payload = b"\xd0\xcf\x11\xe0" + f"NIF {TECNICO.nif}".encode("utf-16-le") + b"\x00" * 32
    add_members(path, {"word/embeddings/oleObject1.bin": payload})


def ooxml_project(root: Path, *, other_requerente: bool = False) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    ficha_eletrotecnica(root / "FE_R9.xlsm", OTHER_REQUERENTE if other_requerente else PROMOTOR)
    identificacao(root / f"Identificacao_{PROMOTOR.name.replace(' ', '_')}.docx")
    termo(root / "Termo_responsabilidade.docx")
    mdj(root / "R9_MDJ_PE_ELE_V0.docx")
    tabela_calculo(root / "Tabela de Calculo.xlsx")
    return root
