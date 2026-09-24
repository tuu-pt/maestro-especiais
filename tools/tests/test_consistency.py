"""A value replaced in one file is replaced in every file, as the verification expects."""

from pathlib import Path

import anonymize
import docx
import pytest

import anonymizer.engine as engine_module
from anonymizer.engine import Allowlist, TextAnonymizer, value_hash
from anonymizer.pseudonyms import PseudonymMap
from anonymizer.textnorm import fold_simple

DETECTED = "O edifício situa-se na Rua das Camélias, 12."  # found by pattern only
UNLABELLED = "moradia na rua das camélias, 12, com dois pisos"  # the pattern does not find it


def write_docx(path: Path, *paragraphs: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(str(path))


def text_of(path: Path) -> str:
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def cli(private: Path, fixtures: Path, *codes: str) -> int:
    return anonymize.main(
        [*codes, "--private-root", str(private), "--fixtures-root", str(fixtures)]
    )


def test_a_value_found_in_a_later_file_is_replaced_in_an_earlier_one(tmp_path: Path) -> None:
    private, fixtures = tmp_path / "private", tmp_path / "fixtures"
    write_docx(private / "R9" / "A_capa.docx", UNLABELLED)  # read before the labelled one
    write_docx(private / "R9" / "B_memoria.docx", DETECTED)

    assert cli(private, fixtures, "R9") == 0

    for name in ("A_capa.docx", "B_memoria.docx"):
        assert "camélias" not in text_of(fixtures / "R9" / name).casefold()


def test_a_value_of_another_project_is_replaced_too(tmp_path: Path) -> None:
    private, fixtures = tmp_path / "private", tmp_path / "fixtures"
    write_docx(private / "R8" / "memoria.docx", DETECTED)
    write_docx(private / "R9" / "capa.docx", UNLABELLED)

    assert cli(private, fixtures, "R8", "R9") == 0
    assert "camélias" not in text_of(fixtures / "R9" / "capa.docx").casefold()


def test_emails_in_a_binary_are_not_reported() -> None:
    data = b"\x00Attribute VB_Name\x00 autor@exemplo.pt \x00"
    assert TextAnonymizer(PseudonymMap()).binary_hits(data) == []


def test_an_allowed_email_in_a_binary_is_not_an_error(monkeypatch: pytest.MonkeyPatch) -> None:
    # Only relevant if emails are anonymized again.
    monkeypatch.setattr(engine_module, "KEPT_KINDS", frozenset())
    data = b"\x00Attribute VB_Name\x00 suporte.modelo@dgeg.gov.pt \x00"
    email = "suporte.modelo@dgeg.gov.pt"
    by_value = Allowlist.from_items([("email", email)])
    by_hash = Allowlist(hashes={f"email:{value_hash(fold_simple(email))}"})  # as in the CI check

    assert TextAnonymizer(PseudonymMap()).binary_hits(data) == ["email"]
    assert TextAnonymizer(PseudonymMap(), allowlist=by_value).binary_hits(data) == []
    assert TextAnonymizer(PseudonymMap(), allowlist=by_hash).binary_hits(data) == []
