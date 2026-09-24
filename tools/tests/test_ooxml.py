import re
import zipfile
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest
import synthetic as s
from runner import Run, Runner

from anonymizer.ooxml import splice
from anonymizer.textnorm import digits_only

# ---------------------------------------------------------------- independent readers


def raw_texts(path: Path) -> list[tuple[str, str]]:
    """Every member of the package, decoded, except image pixel data."""
    out = []
    with zipfile.ZipFile(path) as package:
        for name in package.namelist():
            data = package.read(name)
            if "/media/" in name:
                continue
            if name.lower().endswith((".xlsx", ".xlsm", ".docx")):
                inner = path.parent / f"_inner_{Path(name).name}"
                inner.write_bytes(data)
                out += [(f"{name}›{n}", t) for n, t in raw_texts(inner)]
                inner.unlink()
                continue
            codec = "utf-8" if name.endswith((".xml", ".rels", ".vml")) else "latin-1"
            out.append((name, data.decode(codec, "ignore")))
            if codec == "latin-1":
                out.append((name + "#utf16", data.decode("utf-16-le", "ignore")))
    return out


def docx_texts(path: Path) -> list[str]:
    d = docx.Document(str(path))
    texts = [p.text for p in d.paragraphs]
    for table in d.tables:
        texts += [cell.text for row in table.rows for cell in row.cells]
    for section in d.sections:
        texts += [p.text for p in section.header.paragraphs + section.footer.paragraphs]
    props = d.core_properties
    texts += [props.author or "", props.last_modified_by or ""]
    texts += [c.text for c in d.comments] + [c.author for c in d.comments]
    return texts


def xlsx_texts(path: Path) -> list[str]:
    texts = []
    for data_only in (False, True):
        wb = openpyxl.load_workbook(path, data_only=data_only)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    texts.append("" if cell.value is None else str(cell.value))
                    if cell.comment:
                        texts += [cell.comment.text, cell.comment.author or ""]
            if ws.oddHeader is not None:
                texts.append(ws.oddHeader.center.text or "")
        texts += [wb.properties.creator or "", wb.properties.lastModifiedBy or ""]
    return texts


def assert_no_pii(run: Run) -> None:
    for path in run.fixtures.rglob("*"):
        if not path.is_file():
            continue
        assert s.find_pii(path.name) == [], path.name
        texts = [t for _, t in raw_texts(path)] if path.suffix != ".json" else []
        if path.suffix == ".docx":
            texts += docx_texts(path)
        if path.suffix in {".xlsx", ".xlsm"}:
            texts += xlsx_texts(path)
        for text in texts:
            assert s.find_pii(text) == [], (path.name, s.find_pii(text))


# ---------------------------------------------------------------- tests


@pytest.fixture
def ooxml_run(run_anonymizer: Runner) -> Run:
    return run_anonymizer(s.ooxml_project)


def test_project_is_promoted_without_errors(ooxml_run: Run) -> None:
    assert ooxml_run.outcome.errors == []
    assert ooxml_run.outcome.promoted
    assert len(list(ooxml_run.fixtures.iterdir())) == 5


def test_no_personal_data_survives_anywhere(ooxml_run: Run) -> None:
    assert_no_pii(ooxml_run)


def test_file_names_are_anonymized(ooxml_run: Run) -> None:
    names = [p.name for p in ooxml_run.fixtures.iterdir()]
    assert not any("Almeida" in n for n in names)
    assert any(n.startswith("Identificacao_") and n.endswith(".docx") for n in names)


def test_emails_are_kept(ooxml_run: Run) -> None:
    ws = openpyxl.load_workbook(ooxml_run.output("FE_"), data_only=True).active
    assert ws is not None
    assert ws["C8"].value == s.PROMOTOR.email


def test_technical_content_is_untouched(ooxml_run: Run) -> None:
    text = " ".join(docx_texts(ooxml_run.output("MDJ")))
    for control in ("34,5 kVA", "H07V-U", "XZ1(frt,zh)", "Portaria n.º 949-A/2006", "IP65"):
        assert control in text
    ws = openpyxl.load_workbook(ooxml_run.output("FE_"), data_only=True).active
    assert ws is not None
    assert ws["Q15"].value == "Cedofeita" and ws["G16"].value == "Porto"
    assert ws["J6"].value == "Moradia unifamiliar"


def test_xlsm_keeps_vba_version_and_cached_formula_values(ooxml_run: Run) -> None:
    fe = ooxml_run.output("FE_")
    with zipfile.ZipFile(fe) as package:
        assert package.read("xl/vbaProject.bin") == s.VBA_STUB
    ws = openpyxl.load_workbook(fe, data_only=True).active
    assert ws is not None
    assert ws["R45"].value == "FE_v.20190222"
    assert ws["P29"].value == 34.5
    assert ws["S29"].value == 41.4  # cached value of =P29*1.2
    formulas = openpyxl.load_workbook(fe).active
    assert formulas is not None and formulas["S29"].value == "=P29*1.2"


def test_calculation_table_keeps_values_for_cal01(ooxml_run: Run) -> None:
    ws = openpyxl.load_workbook(ooxml_run.output("Tabela"), data_only=True).active
    assert ws is not None
    assert ws["H3"].value == 504 and ws["I3"].value == 503.4
    assert ws["J3"].value == "XZ1(frt,zh) 4x16"


def test_pseudonyms_are_consistent_across_files(ooxml_run: Run) -> None:
    fe = openpyxl.load_workbook(ooxml_run.output("FE_"), data_only=True).active
    assert fe is not None
    ident = docx.Document(str(ooxml_run.output("Identificacao")))
    promotor = {r.cells[0].text: r.cells[1].text for r in ident.tables[0].rows}

    assert fe["C5"].value == promotor["Nome"]
    assert str(fe["Q5"].value) == digits_only(promotor["NIF"])
    assert promotor["NIF"].count(" ") == 2  # format kept
    mdj = " ".join(docx_texts(ooxml_run.output("MDJ")))
    assert str(fe["C5"].value).upper() in mdj


def test_different_oet_numbers_stay_different_c3(ooxml_run: Run) -> None:
    ident = docx.Document(str(ooxml_run.output("Identificacao")))
    oet_forms = ident.tables[1].cell(1, 2).text
    mdj = " ".join(docx_texts(ooxml_run.output("MDJ")))
    oet_mdj = re.search(r"OET n\.º (\d+)", mdj)

    assert oet_mdj is not None
    assert oet_forms != oet_mdj.group(1)
    assert oet_forms.startswith("0") and oet_mdj.group(1).startswith("0")


def test_form_images_are_blanked_and_photo_metadata_removed(ooxml_run: Run) -> None:
    with zipfile.ZipFile(ooxml_run.output("Identificacao")) as package:
        media = [n for n in package.namelist() if n.startswith("word/media/")]
        image = pymupdf.Pixmap(package.read(media[0]))
    assert set(image.samples) == {255}
    with zipfile.ZipFile(ooxml_run.output("MDJ")) as package:
        photo = package.read(next(n for n in package.namelist() if n.startswith("word/media/")))
        assert b"Exif" not in photo and b"Ferreira" not in photo
        assert not any(n.startswith("docProps/thumbnail") for n in package.namelist())


def test_images_outside_forms_ask_for_visual_review(ooxml_run: Run) -> None:
    mdj = next(f for f in ooxml_run.outcome.files if "MDJ" in f.output)
    assert "images_present" in {x.code for x in mdj.findings}


def test_pii_inside_an_ole_object_blocks_the_project(run_anonymizer: Runner) -> None:
    def build(root: Path) -> None:
        s.ooxml_project(root)
        s.add_ole_with_pii(root / "R9_MDJ_PE_ELE_V0.docx")

    result = run_anonymizer(build)

    assert not result.outcome.promoted
    assert not result.fixtures.exists()
    codes = {(x.code, x.kind) for _, x in result.outcome.errors}
    assert ("pii_in_binary", "nif") in codes


def test_existing_fixtures_stay_untouched_when_a_run_fails(
    run_anonymizer: Runner, tmp_path: Path
) -> None:
    sentinel = tmp_path / "fixtures" / "R9" / "previous.txt"
    sentinel.parent.mkdir(parents=True)
    sentinel.write_text("fixtures anteriores")

    def build(root: Path) -> None:
        s.ooxml_project(root)
        s.add_ole_with_pii(root / "R9_MDJ_PE_ELE_V0.docx")

    run_anonymizer(build)

    assert sentinel.read_text() == "fixtures anteriores"


def test_unknown_ficha_version_is_reported(run_anonymizer: Runner) -> None:
    def build(root: Path) -> None:
        root.mkdir(parents=True)
        path = s.ficha_eletrotecnica(root / "FE.xlsm")
        s.replace_in_member(path, "xl/worksheets/sheet1.xml", b"FE_v.20190222", b"FE_v.20250101")

    result = run_anonymizer(build)
    codes = {x.code for f in result.outcome.files for x in f.findings}
    assert "unknown_fe_version" in codes


def test_splice_keeps_run_boundaries() -> None:
    segments = ["Requerente: ", "Jo", "ão Pedro ", "Almeida", "."]
    joined = "".join(segments)
    start = joined.index("João")
    end = start + len("João Pedro Almeida")

    out = splice(segments, [(start, end, "Ana Exemplo")])

    assert out == ["Requerente: ", "Ana Exemplo", "", "", "."]
