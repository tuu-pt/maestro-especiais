import pdfplumber
import pymupdf
import pytest
import synthetic as s
import xlrd
from runner import Run, Runner


@pytest.fixture
def full_run(run_anonymizer: Runner) -> Run:
    return run_anonymizer(s.full_project)


def codes(run: Run, name_part: str) -> set[str]:
    return {x.code for f in run.outcome.files if name_part in f.output for x in f.findings}


# ---------------------------------------------------------------- whole project


def test_full_project_is_promoted(full_run: Run) -> None:
    assert full_run.outcome.errors == []
    assert full_run.outcome.promoted


def test_unsupported_corrupted_and_encrypted_files_are_never_copied(full_run: Run) -> None:
    names = {p.name for p in full_run.fixtures.rglob("*") if p.is_file()}
    assert not {"R9_EL001.dwg", "assinado.pdf", "protegido.pdf"} & names
    assert "unsupported" in codes(full_run, ".dwg")
    assert "unreadable" in codes(full_run, "assinado.pdf")
    assert "encrypted" in codes(full_run, "protegido.pdf")


# ---------------------------------------------------------------- .xls


def test_xls_has_no_personal_data(full_run: Run) -> None:
    book = xlrd.open_workbook(str(full_run.output("09-Folha")))
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            for c in range(sheet.ncols):
                value = sheet.cell_value(r, c)
                text = (
                    str(int(value))
                    if isinstance(value, float) and value.is_integer()
                    else str(value)
                )
                assert s.find_pii(text) == [], (sheet.name, r, c)
    assert s.find_pii(full_run.output("09-Folha").read_bytes().decode("latin-1")) == []


def test_xls_keeps_values_sheets_and_styles(full_run: Run) -> None:
    book = xlrd.open_workbook(str(full_run.output("09-Folha")), formatting_info=True)
    assert book.sheet_names() == ["IB", "condutores", "tensao", "proteccao"]
    assert book.sheet_by_name("IB").cell_value(6, 7) == 34.5  # IB!H7
    assert book.sheet_by_name("proteccao").cell_value(8, 9) == 504.0  # proteccao!J9
    cell = book.sheet_by_name("IB").cell(1, 0)
    assert book.font_list[book.xf_list[cell.xf_index].font_index].bold
    assert "formulas_to_values" in codes(full_run, "09-Folha")


def test_xls_nif_matches_the_other_files(full_run: Run) -> None:
    import docx

    ident = docx.Document(str(full_run.output("Identificacao")))
    nif = {r.cells[0].text: r.cells[1].text for r in ident.tables[0].rows}["NIF"]
    book = xlrd.open_workbook(str(full_run.output("09-Folha")))
    assert str(int(book.sheet_by_name("IB").cell_value(2, 1))) == nif.replace(" ", "")


# ---------------------------------------------------------------- .pdf


def test_pdf_text_and_metadata_have_no_personal_data(full_run: Run) -> None:
    path = full_run.output("R9_EL_PE_V0.pdf")
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert len(pdf.pages) == 3
        assert not any(pdf.metadata.get(k) for k in ("Author", "Title"))
    assert s.find_pii(text) == []
    for control in ("EL001 Planta de localização", "EL003", "34,5 kVA", "L7 Aplique IP65"):
        assert control in text


def test_pdf_hidden_layers_are_clean(full_run: Run) -> None:
    doc = pymupdf.open(str(full_run.output("R9_EL_PE_V0.pdf")))
    assert doc.embfile_count() == 0
    assert not doc.get_xml_metadata()
    assert all(not list(page.annots() or []) for page in doc)
    widget_values = [str(w.field_value) for page in doc for w in page.widgets() or []]
    assert widget_values and all(s.find_pii(v) == [] for v in widget_values)
    assert all(s.find_pii(title) == [] for _, title, _ in doc.get_toc())
    expanded = doc.tobytes(garbage=0, expand=255).decode("latin-1")
    assert s.find_pii(expanded) == []
    for xref, *_ in doc[0].get_images(full=True):
        assert b"Exif" not in doc.extract_image(xref)["image"]


def test_pdf_vector_pages_and_images_ask_for_visual_review(full_run: Run) -> None:
    found = [x for f in full_run.outcome.files if "R9_EL_PE_V0" in f.output for x in f.findings]
    assert ("possible_vector_text", "pág. 3") in {(x.code, x.where) for x in found}
    assert "images_present" in {x.code for x in found}


def test_pdf_value_hidden_outside_the_text_layer_fails_the_run(run_anonymizer: Runner) -> None:
    def build(root: object) -> None:
        from pathlib import Path

        folder = Path(str(root))
        s.ooxml_project(folder)
        path = folder / "carimbo.pdf"
        s.pecas_desenhadas(path, vector_page=False)
        doc = pymupdf.open(str(path))
        doc.xref_set_key(doc[0].xref, "PieceInfo", f"({s.TECNICO.email})")
        doc.saveIncr()
        doc.close()

    result = run_anonymizer(build)

    assert not result.outcome.promoted
    assert ("pii_in_binary", "email") in {(x.code, x.kind) for _, x in result.outcome.errors}
